import sys
import tempfile
import types
import unittest
from email.message import EmailMessage
from pathlib import Path
from unittest.mock import patch


def _install_google_stubs():
    """Make the utility module importable without installing the ADK runtime.
    Self-contained (includes google.genai, unlike some sibling test files)
    so this module works whether run standalone or as part of the full
    suite, regardless of import order."""
    if "google.adk.models" in sys.modules and "google.genai" in sys.modules:
        return

    google = types.ModuleType("google")
    adk = types.ModuleType("google.adk")
    models = types.ModuleType("google.adk.models")
    agents = types.ModuleType("google.adk.agents")
    callback_context = types.ModuleType("google.adk.agents.callback_context")
    tools = types.ModuleType("google.adk.tools")
    tool_context = types.ModuleType("google.adk.tools.tool_context")
    genai = types.ModuleType("google.genai")
    genai_types = types.ModuleType("google.genai.types")

    class LlmRequest:
        pass

    class LlmResponse:
        pass

    class CallbackContext:
        pass

    class ToolContext:
        pass

    models.LlmRequest = LlmRequest
    models.LlmResponse = LlmResponse
    callback_context.CallbackContext = CallbackContext
    tool_context.ToolContext = ToolContext
    agents.callback_context = callback_context
    tools.tool_context = tool_context
    adk.models = models
    adk.agents = agents
    adk.tools = tools
    google.adk = adk
    google.genai = genai
    genai.types = genai_types

    sys.modules.update(
        {
            "google": google,
            "google.adk": adk,
            "google.adk.models": models,
            "google.adk.agents": agents,
            "google.adk.agents.callback_context": callback_context,
            "google.adk.tools": tools,
            "google.adk.tools.tool_context": tool_context,
            "google.genai": genai,
            "google.genai.types": genai_types,
        }
    )


_install_google_stubs()
from process_agents.common import utils  # noqa: E402


class LoadDirectoryContextTests(unittest.TestCase):
    def setUp(self):
        self.sleep_patch = patch.object(utils, "_safe_sleep_from_property")
        self.sleep_patch.start()
        self.log_patch = patch.object(utils, "_log_agent_activity")
        self.log_patch.start()
        self.addCleanup(self.sleep_patch.stop)
        self.addCleanup(self.log_patch.stop)

    def test_missing_directory_returns_not_found(self):
        with tempfile.TemporaryDirectory() as d:
            missing = str(Path(d) / "does-not-exist")
            result = utils.load_directory_context(missing)
        self.assertEqual(result["status"], "NOT_FOUND")

    def test_extracts_txt_and_docx_and_skips_legacy_doc(self):
        import docx

        with tempfile.TemporaryDirectory() as d:
            (Path(d) / "notes.txt").write_text(
                "Vendor onboarding requires KYC checks and a 3-day SLA."
            )

            doc = docx.Document()
            doc.add_paragraph("Stakeholders: Procurement Lead, Compliance Officer.")
            doc.save(str(Path(d) / "stakeholders.docx"))

            (Path(d) / "legacy.doc").write_text("ignored -- old binary format")

            result = utils.load_directory_context(d)

        self.assertEqual(result["status"], "OK")
        self.assertEqual(sorted(result["files_processed"]), ["notes.txt", "stakeholders.docx"])
        self.assertIn("KYC checks", result["combined_text"])
        self.assertIn("Procurement Lead", result["combined_text"])
        self.assertIn("=== notes.txt ===", result["combined_text"])

        skipped = {entry["file"]: entry["reason"] for entry in result["files_skipped"]}
        self.assertIn("legacy.doc", skipped)
        self.assertIn("re-save as .docx", skipped["legacy.doc"])

    def test_docx_table_text_is_extracted_not_silently_dropped(self):
        import docx

        with tempfile.TemporaryDirectory() as d:
            doc = docx.Document()
            doc.add_paragraph("Requirements register:")
            table = doc.add_table(rows=2, cols=2)
            table.rows[0].cells[0].text = "REQ-1"
            table.rows[0].cells[1].text = "Support SSO login"
            table.rows[1].cells[0].text = "REQ-2"
            table.rows[1].cells[1].text = "Encrypt data at rest"
            doc.save(str(Path(d) / "requirements.docx"))

            result = utils.load_directory_context(d)

        self.assertEqual(result["status"], "OK")
        self.assertIn("Support SSO login", result["combined_text"])
        self.assertIn("Encrypt data at rest", result["combined_text"])

    def test_unsupported_extension_is_skipped_not_fatal(self):
        with tempfile.TemporaryDirectory() as d:
            (Path(d) / "readme.txt").write_text("real content")
            (Path(d) / "image.jpg").write_bytes(b"\xff\xd8\xff\xe0not-a-real-jpeg")

            result = utils.load_directory_context(d)

        self.assertEqual(result["status"], "OK")
        self.assertEqual(result["files_processed"], ["readme.txt"])
        skipped_files = [entry["file"] for entry in result["files_skipped"]]
        self.assertIn("image.jpg", skipped_files)

    def test_excel_xlsx_renders_sheet_as_text(self):
        import pandas as pd

        with tempfile.TemporaryDirectory() as d:
            df = pd.DataFrame({"Step": ["Receive PO", "Validate budget"], "Owner": ["Ops", "Finance"]})
            xlsx_path = Path(d) / "steps.xlsx"
            df.to_excel(xlsx_path, index=False, sheet_name="Process")

            result = utils.load_directory_context(d)

        self.assertEqual(result["status"], "OK")
        self.assertEqual(result["files_processed"], ["steps.xlsx"])
        self.assertIn("Receive PO", result["combined_text"])
        self.assertIn("Validate budget", result["combined_text"])
        self.assertIn("[Sheet: Process]", result["combined_text"])

    def test_eml_extraction_returns_subject_and_body(self):
        with tempfile.TemporaryDirectory() as d:
            msg = EmailMessage()
            msg["Subject"] = "Q3 onboarding requirements"
            msg["From"] = "ops@example.com"
            msg["To"] = "architecture@example.com"
            msg.set_content("All new vendors must complete KYC within 5 business days.")

            eml_path = Path(d) / "requirements.eml"
            eml_path.write_bytes(bytes(msg))

            result = utils.load_directory_context(d)

        self.assertEqual(result["status"], "OK")
        self.assertEqual(result["files_processed"], ["requirements.eml"])
        self.assertIn("Q3 onboarding requirements", result["combined_text"])
        self.assertIn("KYC within 5 business days", result["combined_text"])

    def test_blank_pdf_has_no_extractable_text_and_is_skipped(self):
        from pypdf import PdfWriter

        with tempfile.TemporaryDirectory() as d:
            writer = PdfWriter()
            writer.add_blank_page(width=200, height=200)
            pdf_path = Path(d) / "blank.pdf"
            with open(pdf_path, "wb") as f:
                writer.write(f)

            result = utils.load_directory_context(d)

        self.assertEqual(result["status"], "OK")
        self.assertEqual(result["files_processed"], [])
        skipped = {entry["file"]: entry["reason"] for entry in result["files_skipped"]}
        self.assertEqual(skipped.get("blank.pdf"), "no extractable text")

    def test_msg_extension_is_wired_to_its_extractor(self):
        # A full positive-content test would require fabricating a valid
        # OLE-format .msg file, which extract_msg (read-only) provides no
        # authoring API for. Covering the wiring and the graceful-failure
        # path (a malformed .msg must be skipped, not crash the directory
        # read) is the practical level of coverage here.
        self.assertIs(utils._DIRECTORY_CONTEXT_EXTRACTORS[".msg"], utils._extract_msg_file)

        with tempfile.TemporaryDirectory() as d:
            (Path(d) / "broken.msg").write_bytes(b"not a real OLE msg file")
            result = utils.load_directory_context(d)

        self.assertEqual(result["status"], "OK")
        self.assertEqual(result["files_processed"], [])
        skipped = {entry["file"]: entry["reason"] for entry in result["files_skipped"]}
        self.assertIn("broken.msg", skipped)
        self.assertIn("failed to read", skipped["broken.msg"])

    def test_combined_text_is_truncated_past_the_configured_limit(self):
        with tempfile.TemporaryDirectory() as d:
            (Path(d) / "big.txt").write_text("x" * 500)

            with patch.object(utils, "getProperty", return_value=100):
                result = utils.load_directory_context(d)

        self.assertEqual(result["status"], "OK")
        self.assertTrue(result.get("truncated"))
        self.assertIn("truncated at 100 characters", result["combined_text"])


if __name__ == "__main__":
    unittest.main()
