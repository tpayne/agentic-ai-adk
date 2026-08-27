from pathlib import Path
from typing import Dict, Any, List
import copy
import json

from google.adk.agents.llm_agent import Agent
from google.adk.core import ToolContext

import docx
import openpyxl
from PyPDF2 import PdfReader


# ---------- RFP JSON schema skeleton ----------

EMPTY_RFP_JSON: Dict[str, Any] = {
    "rfp_metadata": {
        "rfp_title": "",
        "rfp_reference_id": "",
        "issuing_organization": "",
        "issue_date": "",
        "due_date": "",
        "contact_person": "",
        "contact_email": "",
        "currency": "",
        "language": "",
        "source_files": []
    },
    "client_context": {
        "organization_overview": "",
        "business_drivers": "",
        "strategic_objectives": "",
        "current_state_summary": "",
        "future_state_vision": ""
    },
    "engagement_scope": {
        "services_in_scope": [],
        "services_out_of_scope": [],
        "geographies_in_scope": [],
        "business_units_in_scope": [],
        "constraints": []
    },
    "deliverables": {
        "primary_deliverables": [],
        "interim_deliverables": [],
        "acceptance_criteria": [],
        "reporting_requirements": []
    },
    "requirements": {
        "functional_requirements": [],
        "non_functional_requirements": [],
        "technical_requirements": [],
        "integration_requirements": [],
        "security_requirements": [],
        "data_privacy_requirements": [],
        "service_level_requirements": [],
        "resource_profile_requirements": [],
        "pricing_requirements": [],
        "commercial_terms_requirements": []
    },
    "iso9001_alignment": {
        "context_of_organization": {
            "clause": "4",
            "summary": "",
            "explicit_references": []
        },
        "leadership": {
            "clause": "5",
            "summary": "",
            "explicit_references": []
        },
        "planning": {
            "clause": "6",
            "summary": "",
            "explicit_references": []
        },
        "support": {
            "clause": "7",
            "summary": "",
            "explicit_references": []
        },
        "operation": {
            "clause": "8",
            "summary": "",
            "explicit_references": []
        },
        "performance_evaluation": {
            "clause": "9",
            "summary": "",
            "explicit_references": []
        },
        "improvement": {
            "clause": "10",
            "summary": "",
            "explicit_references": []
        }
    },
    "evaluation_and_award": {
        "evaluation_criteria": [],
        "weightings_if_stated": [],
        "mandatory_requirements": [],
        "nice_to_have_requirements": [],
        "shortlisting_process": "",
        "presentation_or_demo_requirements": "",
        "contract_award_criteria": ""
    },
    "timeline_and_milestones": {
        "key_dates": [],
        "project_start_target": "",
        "project_end_target": "",
        "milestones_if_stated": []
    },
    "submission_instructions": {
        "submission_format": "",
        "submission_channel": "",
        "number_of_copies": "",
        "page_limits": "",
        "section_structure_required": [],
        "questions_deadline": "",
        "clarification_process": ""
    },
    "risks_and_assumptions": {
        "client_stated_risks": [],
        "implied_risks": [],
        "client_assumptions": [],
        "vendor_assumptions_to_confirm": []
    },
    "compliance_and_legal": {
        "regulatory_requirements": [],
        "standards_and_frameworks": [],
        "ip_and_confidentiality_terms": "",
        "liability_and_indemnity_terms": "",
        "termination_conditions": "",
        "other_legal_terms": ""
    },
    "open_questions_for_vendor": [],
    "analysis_notes": {
        "confidence_level": "",
        "gaps_or_ambiguities": [],
        "files_not_parsed": []
    }
}


# ---------- File readers ----------

def _read_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_docx(path: Path) -> str:
    doc = docx.Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def _read_xlsx(path: Path) -> str:
    wb = openpyxl.load_workbook(str(path), data_only=True)
    chunks: List[str] = []
    for sheet in wb.worksheets:
        chunks.append(f"# SHEET: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            vals = [str(c) for c in row if c is not None]
            if vals:
                chunks.append(" | ".join(vals))
    return "\n".join(chunks)


def _read_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages: List[str] = []
    for i, page in enumerate(reader.pages):
        try:
            pages.append(f"# PAGE {i+1}\n{page.extract_text() or ''}")
        except Exception:
            continue
    return "\n\n".join(pages)


def _read_any(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return _read_txt(path)
    if suffix == ".docx":
        return _read_docx(path)
    if suffix in (".xlsx", ".xlsm", ".xls"):
        return _read_xlsx(path)
    if suffix == ".pdf":
        return _read_pdf(path)
    # Fallback
    return _read_txt(path)


# ---------- Chunking ----------

def _chunk_text(text: str, max_chars: int = 6000) -> List[str]:
    chunks: List[str] = []
    start = 0
    n = len(text)
    while start < n:
        end = min(start + max_chars, n)
        chunks.append(text[start:end])
        start = end
    return chunks


# ---------- JSON merge logic ----------

def _merge_lists(existing: List[Any], new: List[Any]) -> List[Any]:
    seen = set()
    merged: List[Any] = []
    for item in existing + new:
        key = json.dumps(item, sort_keys=True) if isinstance(item, (dict, list)) else str(item)
        if key not in seen:
            seen.add(key)
            merged.append(item)
    return merged


def _merge_rfp_json(base: Dict[str, Any], delta: Dict[str, Any]) -> Dict[str, Any]:
    """
    Merge a sparse delta JSON into the aggregate base JSON.
    - Strings: if delta is non-empty, overwrite.
    - Lists: append-unique.
    - Nested dicts: recurse.
    - Unknown keys are ignored.
    """
    for key, dval in delta.items():
        if key not in base:
            continue
        bval = base[key]

        if isinstance(bval, dict) and isinstance(dval, dict):
            _merge_rfp_json(bval, dval)
        elif isinstance(bval, list) and isinstance(dval, list):
            base[key] = _merge_lists(bval, dval)
        else:
            if dval not in ("", None, [], {}):
                base[key] = dval
    return base


# ---------- Tools ----------

def list_rfp_files(directory: str) -> Dict[str, Any]:
    """
    List supported RFP files in a directory.
    """
    base = Path(directory)
    if not base.exists() or not base.is_dir():
        return {
            "status": "error",
            "message": f"Directory not found: {directory}",
            "files": []
        }

    exts = {".txt", ".docx", ".pdf", ".xlsx", ".xlsm", ".xls"}
    files = [str(p) for p in base.iterdir() if p.is_file() and p.suffix.lower() in exts]

    return {
        "status": "success",
        "directory": str(base),
        "files": files
    }


def analyze_rfp_directory_streaming(directory: str, context: ToolContext | None = None) -> Dict[str, Any]:
    """
    Pass 1: streaming-style analysis (NO iso9001_alignment here).
    - Iterate files
    - Read each file
    - Chunk text
    - For each chunk, call LLM to produce a sparse JSON delta
    - Merge deltas into a single aggregate JSON
    """
    listing = list_rfp_files(directory)
    if listing.get("status") != "success":
        return listing

    if context is None or context.llm is None:
        return {
            "status": "error",
            "message": "No LLM available in ToolContext."
        }

    llm = context.llm
    aggregate: Dict[str, Any] = copy.deepcopy(EMPTY_RFP_JSON)
    files = listing["files"]
    aggregate["rfp_metadata"]["source_files"] = files

    # Schema shape WITHOUT iso9001_alignment for this pass
    schema_for_chunks = copy.deepcopy(EMPTY_RFP_JSON)
    schema_for_chunks.pop("iso9001_alignment", None)
    schema_description = json.dumps(schema_for_chunks, indent=2)

    for path_str in files:
        path = Path(path_str)
        try:
            text = _read_any(path)
        except Exception as e:
            aggregate["analysis_notes"]["files_not_parsed"].append(
                {"file": str(path), "error": str(e)}
            )
            continue

        chunks = _chunk_text(text, max_chars=6000)

        for idx, chunk in enumerate(chunks):
            prompt = f"""
You are an expert consulting RFP analyst.

You are given a CHUNK of an RFP-related document.
Your job is to extract ONLY the information that is clearly supported by this chunk
and map it into a PARTIAL JSON object that follows this schema:

<SCHEMA SHAPE>
{schema_description}
</SCHEMA SHAPE>

IMPORTANT:
- Do NOT include the field "iso9001_alignment" in your output at all.
- You are only populating factual RFP content (requirements, scope, deliverables, risks, etc.).

Rules:
- Output a SINGLE JSON object.
- Include ONLY fields you can populate with evidence from THIS CHUNK.
- For fields you cannot populate from this chunk, OMIT them entirely from the JSON.
- Do NOT invent specific facts (names, dates, amounts, etc.).
- You may add to list fields (e.g. requirements, risks, evaluation criteria) when they are clearly present.

Context:
- File: {path.name}
- Chunk index: {idx+1} of {len(chunks)}

<CHUNK>
{chunk}
</CHUNK>

Return ONLY the JSON object, no explanation.
"""
            resp = llm(prompt)

            try:
                delta = json.loads(resp)
            except Exception:
                aggregate["analysis_notes"]["gaps_or_ambiguities"].append(
                    {
                        "file": str(path),
                        "chunk_index": idx,
                        "issue": "Non-JSON or malformed JSON from model."
                    }
                )
                continue

            _merge_rfp_json(aggregate, delta)

    # crude confidence heuristic – you can refine this
    aggregate["analysis_notes"]["confidence_level"] = "medium"

    return {
        "status": "success",
        "directory": directory,
        "rfp_analysis": aggregate
    }


def generate_iso9001_alignment(rfp_json: Dict[str, Any], context: ToolContext | None = None) -> Dict[str, Any]:
    """
    Pass 2: given the merged RFP JSON (without ISO alignment),
    generate ONLY the iso9001_alignment section in one holistic call.
    """
    if context is None or context.llm is None:
        return {
            "status": "error",
            "message": "No LLM available in ToolContext."
        }

    llm = context.llm

    prompt = f"""
You are an ISO 9001:2015 expert.

You are given a structured JSON representation of an RFP's requirements and context.
Your task is to produce ONLY the "iso9001_alignment" section, mapping the RFP content
to the relevant ISO 9001 clauses.

<RFP_JSON>
{json.dumps(rfp_json, indent=2)}
</RFP_JSON>

Return a JSON object with EXACTLY this shape:

{{
  "iso9001_alignment": {{
    "context_of_organization": {{"clause": "4", "summary": "", "explicit_references": []}},
    "leadership": {{"clause": "5", "summary": "", "explicit_references": []}},
    "planning": {{"clause": "6", "summary": "", "explicit_references": []}},
    "support": {{"clause": "7", "summary": "", "explicit_references": []}},
    "operation": {{"clause": "8", "summary": "", "explicit_references": []}},
    "performance_evaluation": {{"clause": "9", "summary": "", "explicit_references": []}},
    "improvement": {{"clause": "10", "summary": "", "explicit_references": []}}
  }}
}}

Guidance:
- "summary" should briefly describe how the RFP content relates to that clause.
- "explicit_references" should list any direct mentions or strong implied references
  (e.g. quality management, continuous improvement, performance monitoring, etc.).
- If there is no clear linkage for a clause, leave its "summary" as an empty string
  and "explicit_references" as an empty list.

Return ONLY the JSON object, no explanation.
"""
    resp = llm(prompt)

    try:
        iso_obj = json.loads(resp)
    except Exception:
        return {
            "status": "error",
            "message": "Malformed JSON from ISO alignment model call."
        }

    return {
        "status": "success",
        "iso9001_alignment": iso_obj.get("iso9001_alignment", {})
    }


def analyze_rfp_directory_full(directory: str, context: ToolContext | None = None) -> Dict[str, Any]:
    """
    Public entrypoint:
    - Pass 1: streaming extraction (no ISO)
    - Pass 2: holistic ISO 9001 alignment
    - Merge and return final JSON
    """
    # Pass 1
    streaming_result = analyze_rfp_directory_streaming(directory, context)
    if streaming_result.get("status") != "success":
        return streaming_result

    rfp_json = streaming_result["rfp_analysis"]

    # Pass 2
    iso_result = generate_iso9001_alignment(rfp_json, context)
    if iso_result.get("status") != "success":
        # keep non-fatal; just record the issue
        rfp_json["analysis_notes"]["gaps_or_ambiguities"].append(
            {
                "file": None,
                "chunk_index": None,
                "issue": f"ISO alignment failed: {iso_result.get('message')}"
            }
        )
    else:
        rfp_json["iso9001_alignment"] = iso_result["iso9001_alignment"]

    return {
        "status": "success",
        "directory": directory,
        "rfp_analysis": rfp_json
    }


# ---------- Root agent ----------

root_agent = Agent(
    model="gemini-3-flash-preview",
    name="rfp_request_analysis_agent",
    description=(
        "Two-pass RFP analysis agent. "
        "Pass 1: streaming extraction of RFP structure and requirements. "
        "Pass 2: holistic ISO 9001 alignment over the merged JSON."
    ),
    instruction=(
        "You are a consulting RFP analysis agent. "
        "When the user provides a directory path, call analyze_rfp_directory_full "
        "to produce the final JSON. "
        "Do not attempt to read or summarize files yourself; always use the tools."
    ),
    tools=[list_rfp_files, analyze_rfp_directory_full],
)
