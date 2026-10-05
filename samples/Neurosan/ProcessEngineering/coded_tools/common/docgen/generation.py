"""Generates a structured, ISO-formatted Word document from the master
process or design-document JSON. Ported from the ADK sample's
process_agents/common/doc_generation_agent.py -- both
`_build_process_document` and `_build_design_document` (+
`_add_part_divider`/`_renumber_design_document_headings`).

Also dropped: the ADK original's getProperty("modelSleep")-based random
jitter sleep at the top of this function, which existed purely to spread
out this app's own ADK-side LLM rate-limit retries -- not relevant here,
since neuro-san's own LLM call path isn't this port's concern.
"""

import os
import re
import json
import traceback
import logging
from datetime import datetime
from typing import Optional

import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from coded_tools.common import paths
from coded_tools.common.filenames import safe_filename_component
from coded_tools.common.process_json import PROCESS_JSON_FILENAME, detect_schema_type_from_disk
from coded_tools.common.process_json import DESIGN_JSON_FILENAME
from coded_tools.common.docgen.themes import apply_theme
from coded_tools.common.docgen.structure import (
    _add_version_history_table,
    _add_table_of_contents,
    add_iso_page_break,
)
from coded_tools.common.docgen.content import (
    _add_overview_section,
    _add_design_overview_section,
    _add_stakeholders_section,
    _add_process_steps_section,
    _render_generic_value,
)
from coded_tools.common.docgen.technical import (
    _add_tools_section_from_summary,
    _add_metrics_section,
    _add_system_requirements,
    _add_flowchart_section,
    _add_simulation_report,
)
from coded_tools.common.docgen.governance import (
    _add_governance_requirements_section,
    _add_risks_and_controls_section,
    _add_process_triggers_section,
    _add_process_end_conditions_section,
    _add_change_management_section,
    _add_continuous_improvement_section,
    _add_appendix_from_json,
    _add_additional_data_section,
    _add_glossary,
    _add_critical_success_factors_section,
    _add_critical_failure_factors_section,
    _add_reporting_and_analytics,
)
from coded_tools.common.docgen.design_sections import (
    _add_requirements_section,
    _add_system_context_section,
    _add_quality_attributes_section,
    _add_compliance_and_standards_section,
    _add_risk_register_section,
    _add_architecture_analysis,
    _add_low_level_design_section,
    _add_glossary_and_references_section,
)

logger = logging.getLogger("ProcessArchitect.DocGen")

SIM_RESULTS_FILENAME = "simulation_results.json"


# ============================================================
# LOAD SUBPROCESSES
# ============================================================

def _load_subprocesses() -> dict:
    """Loads all subprocess JSON files from output/subprocesses/."""
    logger.debug("Loading subprocess JSON files...")
    subprocess_dir = os.path.join(paths.OUTPUT_DIR, "subprocesses")
    subprocesses = {}

    if not os.path.isdir(subprocess_dir):
        return subprocesses

    for filename in os.listdir(subprocess_dir):
        if not filename.endswith(".json"):
            continue

        path = os.path.join(subprocess_dir, filename)
        try:
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)

            parent = data.get("parent_step_name") or data.get("step_name")
            if parent:
                data["step_name"] = parent
                subprocesses[parent] = data

        except Exception:
            logger.exception(f"Failed to load subprocess file: {path}")

    logger.debug("Subprocesses loaded.")
    return subprocesses


# ============================================================
# GLOBAL DOCUMENT STYLE SETUP (fallback when no theme is applied)
# ============================================================

def _apply_global_styles(doc: docx.Document):
    """Apply ISO-style global typography and spacing."""
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.15

    h1 = styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(12)

    h2 = styles["Heading 2"]
    h2.font.name = "Calibri"
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)

    h3 = styles["Heading 3"]
    h3.font.name = "Calibri"
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.paragraph_format.space_before = Pt(6)
    h3.paragraph_format.space_after = Pt(3)

    h4 = styles["Heading 4"]
    h4.font.name = "Calibri"
    h4.font.size = Pt(11)
    h4.font.bold = True
    h4.paragraph_format.space_before = Pt(6)
    h4.paragraph_format.space_after = Pt(3)

    for section in doc.sections:
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)


# ============================================================
# MAIN DOCUMENT GENERATION
# ============================================================

def create_standard_doc_from_file(process_name: str = "Process", schema_type: Optional[str] = None) -> str:
    """
    Generate a structured, ISO-formatted Word document from the master
    JSON (output/process_data.json or output/design_data.json).
    `schema_type` defaults to auto-detecting from which master file exists
    on disk (detect_schema_type_from_disk).
    """
    logger.debug(f"Creating document for process: {process_name}...")

    try:
        resolved_type = schema_type or detect_schema_type_from_disk()
        data_filename = DESIGN_JSON_FILENAME if resolved_type == "design" else PROCESS_JSON_FILENAME

        data_path = paths.output_path(data_filename)
        with open(data_path, "r", encoding="utf-8") as f:
            raw_data = json.load(f)

        data = raw_data.get("process_design") if isinstance(raw_data, dict) and "process_design" in raw_data else raw_data

        doc = docx.Document()
        apply_theme(doc)

        if resolved_type == "design":
            out_path = _build_design_document(doc, data, process_name)
        else:
            out_path = _build_process_document(doc, data, process_name)

        return f"SUCCESS: Professional document saved at {out_path}"

    except Exception as e:
        traceback.print_exc()
        return f"ERROR: {str(e)}"


def _build_process_document(doc: docx.Document, data: dict, process_name: str) -> str:
    """Builds the full ISO-formatted process document into `doc` and saves
    it. Returns the saved file path."""
    # Attach subprocesses
    subprocesses = _load_subprocesses()
    for step in data.get("process_steps", []):
        if isinstance(step, dict):
            name = step.get("step_name")
            if name in subprocesses:
                step["subprocess"] = subprocesses[name]

    # Extract metadata
    name = str(data.get("process_name", process_name))
    version = str(data.get("version", "1.0"))
    sector = data.get("industry_sector", data.get("business_unit", "N/A"))

    stakeholders = data.get("stakeholders")
    process_steps = data.get("process_steps")
    tools_summary = data.get("tools_summary")
    critical_success_factors = data.get("critical_success_factors")
    critical_failure_factors = data.get("critical_failure_factors")
    metrics = data.get("metrics") or data.get("success_metrics")
    reporting_and_analytics = data.get("reporting_and_analytics")
    system_requirements = data.get("system_requirements")
    appendix = data.get("appendix") if isinstance(data.get("appendix"), dict) else None

    governance_requirements = data.get("governance_requirements")
    process_end_conditions = data.get("process_end_conditions")
    change_management = data.get("change_management")
    process_triggers = data.get("process_triggers")
    continuous_improvement = data.get("continuous_improvement")
    risks_and_controls = data.get("risks_and_controls")

    consumed_keys = {
        "appendix", "assumptions", "business_unit", "change_management",
        "constraints", "continuous_improvement", "critical_failure_factors",
        "critical_success_factors", "description", "governance_requirements",
        "industry_sector", "introduction", "metrics", "process_description",
        "process_end_conditions", "process_name", "process_steps",
        "process_triggers", "reporting_and_analytics", "risks_and_controls",
        "stakeholders", "success_metrics", "system_requirements",
        "tools_summary", "version",
        # "purpose"/"scope" ARE rendered (in _add_overview_section's own
        # "ordered" loop) but were missing here, so they rendered correctly
        # in 1.0 Overview and then ALSO reappeared, duplicated, in Appendix
        # B's leftover-data catch-all below. "process_owner"/"owner"
        # likewise (see _add_overview_section's own fix for why both keys
        # are checked).
        "purpose", "scope", "process_owner", "owner",
    }

    # ---- COVER PAGE ----
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(name)
    run.bold = True
    run.font.size = Pt(28)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Process Model Document").font.size = Pt(18)

    doc.add_paragraph()
    doc.add_paragraph(f"Industry / Domain: {sector}")
    doc.add_paragraph(f"Version: {version}")
    doc.add_paragraph(f"Generated On: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    add_iso_page_break(doc)

    # ---- DOCUMENT CONTROL + TOC ----
    _add_version_history_table(doc, version=version, author="Process Architect")

    add_iso_page_break(doc)
    doc.add_heading("Table of Contents", level=1)
    _add_table_of_contents(doc)

    add_iso_page_break(doc)

    # ---- MAIN SECTIONS ----
    _add_overview_section(doc, data)
    add_iso_page_break(doc)

    _add_stakeholders_section(doc, stakeholders)
    add_iso_page_break(doc)

    if process_steps:
        _add_process_steps_section(doc, process_steps)
        add_iso_page_break(doc)

    if tools_summary:
        _add_tools_section_from_summary(doc, tools_summary)
        add_iso_page_break(doc)

    if isinstance(metrics, list) and metrics:
        _add_metrics_section(doc, metrics)
        add_iso_page_break(doc)

    if isinstance(critical_success_factors, list) and critical_success_factors:
        _add_critical_success_factors_section(doc, critical_success_factors)
        add_iso_page_break(doc)

    if isinstance(critical_failure_factors, list) and critical_failure_factors:
        _add_critical_failure_factors_section(doc, critical_failure_factors)
        add_iso_page_break(doc)

    if reporting_and_analytics:
        _add_reporting_and_analytics(doc, reporting_and_analytics)
        add_iso_page_break(doc)

    if system_requirements:
        _add_system_requirements(doc, system_requirements)
        add_iso_page_break(doc)

    _add_flowchart_section(doc, name)
    add_iso_page_break(doc)

    # Simulation results
    simulation_results = None
    try:
        sim_path = paths.output_path(SIM_RESULTS_FILENAME)
        if os.path.exists(sim_path):
            with open(sim_path, "r", encoding="utf-8") as sf:
                simulation_results = json.load(sf)
    except Exception:
        traceback.print_exc()

    if simulation_results:
        _add_simulation_report(doc, simulation_results)
        add_iso_page_break(doc)

    if governance_requirements:
        _add_governance_requirements_section(doc, governance_requirements)
        add_iso_page_break(doc)

    if risks_and_controls:
        _add_risks_and_controls_section(doc, risks_and_controls)
        add_iso_page_break(doc)

    if process_triggers:
        _add_process_triggers_section(doc, process_triggers)
        add_iso_page_break(doc)

    if process_end_conditions:
        _add_process_end_conditions_section(doc, process_end_conditions)
        add_iso_page_break(doc)

    if change_management:
        _add_change_management_section(doc, change_management)
        add_iso_page_break(doc)

    if continuous_improvement:
        _add_continuous_improvement_section(doc, continuous_improvement)
        add_iso_page_break(doc)

    # Appendices
    if appendix:
        _add_appendix_from_json(doc, appendix)
        consumed_keys.add("appendix")

    add_iso_page_break(doc)
    _add_additional_data_section(doc, data, consumed_keys)

    add_iso_page_break(doc)
    _add_glossary(doc)

    # Save
    out_path = paths.output_path(f"{safe_filename_component(name)}.docx")
    doc.save(out_path)
    return out_path


def _add_part_divider(doc: docx.Document, label: str) -> bool:
    """
    Adds a visually distinct 'Part' divider heading, used to physically
    separate the high-level architecture (HLD) sections of a combined
    design document from the low-level design (LLD) implementation
    detail that follows. Always renders (no data-dependent no-op).
    """
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(220)
    run = p.add_run(label)
    run.bold = True
    run.font.size = Pt(24)
    return True


def _build_design_document(doc: docx.Document, data: dict, process_name: str) -> str:
    """
    Builds a foundational ISO-formatted design (HLD/LLD/Combined) document
    into `doc` and saves it. Design-schema counterpart to
    _build_process_document.

    Deliberately foundation-level, not a full bespoke renderer for every
    nested design_document_schema.json section the way the process
    document has hand-crafted, data-shape-specific renderers for every
    field: most sections go through _render_generic_value, which already
    produces proper ISO-formatted Word tables/lists for any dict/
    list-of-dict/list-of-string shape. Sections with a bespoke renderer
    (requirements, system context, quality attributes, compliance and
    standards, risk register, architecture analysis, low-level design,
    glossary and references) live in design_sections.py.
    """
    metadata = data.get("document_metadata") or {}

    name = str(metadata.get("title") or metadata.get("system_name") or process_name)
    # Short system name for diagram hub labels -- "name" above is the full
    # document TITLE (a whole sentence), fine for a cover page but
    # unreadable as a diagram node label. "system_name" is the short form
    # meant for exactly this.
    system_label = str(metadata.get("system_name") or name)
    version = str(metadata.get("version", "1.0"))
    doc_type = str(metadata.get("document_type", "Design"))
    sector = metadata.get("industry_sector", "N/A")

    architecture_description = data.get("architecture_description") or {}
    stakeholders = architecture_description.get("stakeholders")

    requirements = data.get("requirements")
    system_context = data.get("system_context")
    high_level_design = data.get("high_level_design") or {}
    low_level_design = data.get("low_level_design")
    quality_attributes = data.get("quality_attributes")
    compliance_and_standards = data.get("compliance_and_standards")
    risk_register = data.get("risk_register")
    governance_requirements = data.get("governance_requirements")
    continuous_improvement = data.get("continuous_improvement")
    glossary_and_references = data.get("glossary_and_references")
    appendix = data.get("appendix") if isinstance(data.get("appendix"), dict) else None

    consumed_keys = {
        "document_metadata", "business_context", "architecture_description",
        "requirements", "system_context", "high_level_design",
        "low_level_design", "quality_attributes", "compliance_and_standards",
        "risk_register", "governance_requirements", "continuous_improvement",
        "glossary_and_references", "review_and_governance", "appendix",
    }

    # ---- COVER PAGE ----
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(name)
    run.bold = True
    run.font.size = Pt(28)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"{doc_type} Design Document").font.size = Pt(18)

    doc.add_paragraph()
    doc.add_paragraph(f"Industry / Domain: {sector}")
    doc.add_paragraph(f"Version: {version}")
    doc.add_paragraph(f"Generated On: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    add_iso_page_break(doc)

    # ---- DOCUMENT CONTROL + TOC ----
    _add_version_history_table(
        doc, version=version, author="Design Architect",
        description="Initial generated design specification",
    )

    add_iso_page_break(doc)
    doc.add_heading("Table of Contents", level=1)
    _add_table_of_contents(doc)

    add_iso_page_break(doc)

    # ---- MAIN SECTIONS ----
    # Every section below is only page-broken from the NEXT one if it
    # actually rendered something (see `rendered` tracking) -- an
    # unconditional break both after a no-op section and before the next
    # one stacks two explicit page breaks with nothing between them, which
    # Word renders as a genuinely blank page.
    #
    # Section numbers are intentionally left un-normalized here;
    # _renumber_design_document_headings fixes them all in one pass from
    # the final actual set of headings, right before saving.
    #
    # Sections are grouped into two visually divided parts rather than raw
    # JSON key order: Part I is everything a reader needs to approve the
    # design; Part II is everything an implementing engineer needs to
    # build it.

    # ---------------- PART I: HIGH-LEVEL ARCHITECTURE (HLD) ----------------
    _add_part_divider(doc, "PART I: HIGH-LEVEL ARCHITECTURE (HLD)")
    add_iso_page_break(doc)

    rendered = True  # Document Overview always renders unconditionally.
    _add_design_overview_section(doc, data)

    if stakeholders:
        if rendered:
            add_iso_page_break(doc)
        rendered = _add_stakeholders_section(doc, stakeholders, subject_noun="architecture")

    if requirements:
        if rendered:
            add_iso_page_break(doc)
        rendered = _add_requirements_section(doc, requirements, heading="3.0 Requirements")

    if system_context:
        if rendered:
            add_iso_page_break(doc)
        rendered = _add_system_context_section(
            doc, system_context, heading="4.0 System Context", system_name=system_label
        )

    if architecture_description:
        if rendered:
            add_iso_page_break(doc)
        doc.add_heading("5.0 Architecture Description", level=1)
        doc.add_paragraph(
            "The following describes the architecture viewpoints and views "
            "for this design (ISO/IEC/IEEE 42010), and the analysis behind "
            "the recommended approach."
        )
        for key in ("concerns", "viewpoints", "views"):
            value = architecture_description.get(key)
            if value:
                _render_generic_value(doc, value, label=key.replace("_", " ").title(), system_name=system_label)

        architecture_analysis = architecture_description.get("architecture_analysis")
        if architecture_analysis:
            _add_architecture_analysis(doc, architecture_analysis, level=3)

        rendered = True

    # high_level_design's own architectural detail -- components,
    # integration points, deployment topology, security/scalability/
    # availability characteristics, and architecture-level risks. None of
    # these have a bespoke renderer of their own; they go through
    # _render_generic_value (the same generic fallback every other
    # schema-conforming-but-not-specially-handled section in this
    # document uses), labeled so a reader can tell which schema field
    # each came from. Previously high_level_design was only ever added to
    # consumed_keys (to keep it out of Appendix B) and never actually
    # rendered anywhere in the document body at all -- a complete,
    # silent loss of real content (e.g. a real design's actual security
    # architecture prose) confirmed directly against a real generated
    # document.
    if isinstance(high_level_design, dict) and high_level_design:
        hld_sections = [
            ("solution_overview", "Solution Overview"),
            ("architecture_style", "Architecture Style"),
            ("components", "Components"),
            ("integration_points", "Integration Points"),
            ("data_flow_overview", "Data Flow Overview"),
            ("technology_stack", "Technology Stack"),
            ("deployment_topology", "Deployment Topology"),
            ("security_architecture", "Security Architecture"),
            ("scalability_and_performance", "Scalability and Performance"),
            ("availability_and_resilience", "Availability and Resilience"),
            ("risks_and_mitigations", "Architecture-Level Risks and Mitigations"),
        ]
        present = [(key, label) for key, label in hld_sections if high_level_design.get(key)]
        if present:
            if rendered:
                add_iso_page_break(doc)
            doc.add_heading("5.1 Architecture Detail", level=2)
            doc.add_paragraph(
                "The following elaborates the high-level design's own components, "
                "integration points, deployment topology, and cross-cutting "
                "security, scalability, and availability characteristics."
            )
            for key, label in present:
                _render_generic_value(
                    doc, high_level_design[key], label=label, level=3, system_name=system_label
                )
            rendered = True

    # Architecture/component diagram (output/<sanitized name>_flow.png,
    # generated by edge_inference.py's design-schema branch). Called
    # unconditionally -- _add_flowchart_section itself checks whether the
    # file exists and returns False (a silent no-op) if it doesn't.
    if rendered:
        add_iso_page_break(doc)
    rendered = _add_flowchart_section(
        doc, name, heading="6.0 Architecture Flow Diagram",
        caption=(
            "This section provides a high-level visualization of the "
            "architecture and the flow between its major components."
        ),
    )

    if quality_attributes:
        if rendered:
            add_iso_page_break(doc)
        rendered = _add_quality_attributes_section(doc, quality_attributes, heading="7.0 Quality Attributes")

    if compliance_and_standards:
        if rendered:
            add_iso_page_break(doc)
        rendered = _add_compliance_and_standards_section(
            doc, compliance_and_standards, heading="8.0 Compliance and Standards"
        )

    if risk_register:
        if rendered:
            add_iso_page_break(doc)
        rendered = _add_risk_register_section(doc, risk_register, heading="9.0 Risk Register")

    if governance_requirements:
        if rendered:
            add_iso_page_break(doc)
        rendered = _add_governance_requirements_section(
            doc, governance_requirements, subject_noun="architecture"
        )

    if continuous_improvement:
        if rendered:
            add_iso_page_break(doc)
        rendered = _add_continuous_improvement_section(
            doc, continuous_improvement, subject_noun="architecture"
        )

    # review_and_governance (review_history/change_log/approval_workflow)
    # is deliberately NOT rendered -- it's fabricated by the generating
    # LLM with no grounding (invented reviewer names, invented dates) and
    # adds no value. It stays in `consumed_keys` so it doesn't reappear
    # via the Appendix B leftover-data catch-all below.

    # ---------------- PART II: LOW-LEVEL DESIGN (LLD) ----------------
    if rendered:
        add_iso_page_break(doc)
    rendered = _add_part_divider(doc, "PART II: LOW-LEVEL DESIGN (LLD)")

    if low_level_design:
        if rendered:
            add_iso_page_break(doc)
        rendered = _add_low_level_design_section(
            doc, low_level_design, heading="5.0 Low-Level Design", system_name=system_label
        )

    if glossary_and_references:
        if rendered:
            add_iso_page_break(doc)
        rendered = _add_glossary_and_references_section(
            doc, glossary_and_references, heading="11.0 Glossary and References"
        )

    # Appendices
    if appendix:
        if rendered:
            add_iso_page_break(doc)
        rendered = _add_appendix_from_json(doc, appendix, subject_noun="architecture")
        consumed_keys.add("appendix")

    if rendered:
        add_iso_page_break(doc)
    rendered = _add_additional_data_section(doc, data, consumed_keys)

    # _add_glossary is the process document's fallback glossary (four
    # hardcoded process-specific terms) -- not called here at all. A
    # design document with no glossary_and_references data of its own
    # should simply have no glossary section, the same way every diagram/
    # table renderer elsewhere in this pipeline withholds a section rather
    # than fabricating one.

    # Fix up section numbers/appendix letters from the final set of
    # headings that actually got rendered above.
    _renumber_design_document_headings(doc)

    # Save
    out_path = paths.output_path(f"{safe_filename_component(name)}.docx")
    doc.save(out_path)
    return out_path


# ============================================================
# HEADING RENUMBERING (design document)
# ============================================================

_LEADING_NUMBER_RE = re.compile(r"^\d+(?:\.\d+)*\s+")
_APPENDIX_PREFIX_RE = re.compile(r"^Appendix\s+[A-Za-z]:\s*", re.IGNORECASE)
_SUBSECTION_NUMBER_RE = re.compile(r"^\d+\.(\d+)\s+(.*)$")


def _set_paragraph_text(paragraph, new_text: str) -> None:
    """Replaces a heading paragraph's visible text while keeping its
    existing run formatting. doc.add_heading() always creates the heading
    text in a single run, so this only needs to touch runs[0]."""
    if not paragraph.runs:
        paragraph.add_run(new_text)
        return
    paragraph.runs[0].text = new_text
    for extra_run in paragraph.runs[1:]:
        extra_run.text = ""


def _renumber_design_document_headings(doc: docx.Document) -> None:
    """
    Re-derives every top-level section number and appendix letter from
    the FINAL, actual set of Heading 1 paragraphs in the document, in the
    order they were really rendered -- rather than trusting the literal
    "N.0 ..."/"Appendix X: ..." text each section-adding function
    hardcodes when it writes its own heading.

    This is necessary because several section renderers here carry fixed
    numbers that assume every earlier optional section always renders,
    which is not the normal case -- which sections render depends on
    what's actually present in the source JSON. Renumbering after the
    fact sidesteps collisions/gaps: a section that didn't render is
    simply not there to be numbered.

    Heading 2 subsections (e.g. "12.1 Runtime Processing and Sequence
    Flows") are also corrected: their trailing ".Y title" is kept as-is,
    and only the leading section number is swapped for whatever
    `section_no` its enclosing Heading 1 was actually just corrected to.
    """
    section_no = 0
    appendix_idx = 0
    for paragraph in doc.paragraphs:
        if paragraph.style is None:
            continue

        if paragraph.style.name == "Heading 2":
            match = _SUBSECTION_NUMBER_RE.match(paragraph.text)
            if match and section_no:
                subsection_idx, rest = match.groups()
                _set_paragraph_text(paragraph, f"{section_no}.{subsection_idx} {rest}")
            continue

        if paragraph.style.name != "Heading 1":
            continue
        text = paragraph.text

        if _APPENDIX_PREFIX_RE.match(text):
            rest = _APPENDIX_PREFIX_RE.sub("", text).strip()
            letter = chr(ord("A") + appendix_idx)
            appendix_idx += 1
            _set_paragraph_text(paragraph, f"Appendix {letter}: {rest}")
        elif _LEADING_NUMBER_RE.match(text):
            rest = _LEADING_NUMBER_RE.sub("", text).strip()
            section_no += 1
            _set_paragraph_text(paragraph, f"{section_no}.0 {rest}")
        # else: unnumbered heading (Document Control, Table of Contents) -- leave as-is.
