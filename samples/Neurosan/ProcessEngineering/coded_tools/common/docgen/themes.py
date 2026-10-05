"""Named Word-document theme loader/applier. Ported from the ADK sample's
process_agents/common/helpers/themes/loader.py, with the single theme JSON
it shipped (corporate_standard.json) inlined as a dict instead of a sibling
file -- the ADK original's getProperty("theme") config toggle (an
agentapp.properties setting) has no equivalent in this port, so
create_standard_doc_from_file (coded_tools/common/docgen/generation.py)
always applies this one theme rather than selecting between a configured
theme name and the plain-Calibri fallback.
"""

from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

CORPORATE_STANDARD_THEME = {
    "name": "Corporate Standard",
    "fonts": {"heading": "Segoe UI", "body": "Calibri"},
    "colors": {"primary": None, "secondary": "2E75B6", "accent": "5B9BD5"},
    "heading_sizes": {"h1": 18, "h2": 16, "h3": 14, "h4": 12, "h5": 10},
    "body_size": 11,
    "footer_text": "Confidential",
}


def _apply_color(style, rgb_hex) -> None:
    """Apply a hex RGB color to a style if rgb_hex is valid. Null or empty
    values are ignored safely."""
    if not rgb_hex:
        return
    try:
        style.font.color.rgb = RGBColor.from_string(rgb_hex)
    except Exception:
        pass


def apply_theme(doc, theme: dict = None) -> None:
    """
    Apply a theme dict to a python-docx Document. Modifies: Normal style,
    Heading 1-5 styles, Table Grid style, and the first section's footer
    text.
    """
    theme = theme or CORPORATE_STANDARD_THEME
    fonts = theme["fonts"]
    colors = theme["colors"]
    heading_sizes = theme["heading_sizes"]
    body_size = theme["body_size"]

    try:
        normal = doc.styles["Normal"]
        normal.font.name = fonts.get("body", "Calibri")
        normal.font.size = Pt(body_size)
    except KeyError:
        pass

    def style_heading(name: str, size_key: str, color_key: str = None):
        try:
            style = doc.styles[name]
        except KeyError:
            return
        style.font.name = fonts.get("heading", "Segoe UI")
        style.font.size = Pt(heading_sizes[size_key])
        style.font.bold = True
        if color_key:
            _apply_color(style, colors.get(color_key))

    style_heading("Heading 1", "h1", "primary")
    style_heading("Heading 2", "h2", "primary")
    style_heading("Heading 3", "h3", "secondary")
    style_heading("Heading 4", "h4")
    style_heading("Heading 5", "h5")

    try:
        table_style = doc.styles["Table Grid"]
        table_style.font.name = fonts.get("body", "Calibri")
        table_style.font.size = Pt(body_size)
    except KeyError:
        pass

    footer_text = theme.get("footer_text")
    if footer_text:
        section = doc.sections[0]
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.text = footer_text
        p.style = doc.styles["Normal"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
