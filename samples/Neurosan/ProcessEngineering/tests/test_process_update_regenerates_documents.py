"""Regression tests for two real gaps found while documenting this port's
network composition (see codeDoc/README.md):

1. `design_update.hocon`'s front-man correctly ends its update sequence with
   generate_design_flow_diagram/generate_design_document (regenerating the
   Word document and diagram after every update, unconditionally), but
   `process_update.hocon`'s own front-man had no equivalent -- an update to
   an existing process was never followed by document regeneration, only a
   plain-language outcome report, despite the main README's own "Applying
   Updates to Existing Processes & Regenerating the Documentation" sample
   prompts promising exactly that. Fixed by adding the same two tools (and
   their own re-export modules under coded_tools/process_update/, mirroring
   coded_tools/design_update/'s existing pattern) to process_update.hocon.

2. The ADK original's own Update_Design_Pipeline also re-runs subprocess
   regeneration (process_agents/process/update_process_agent.py) before
   rebuilding artifacts on an update, but process_update.hocon had no
   Subprocess_Driver_Agent/Subprocess_Generator_Agent/load_process_steps/
   save_subprocess_flow at all -- updating a process never re-expanded its
   per-step subprocesses, even when the update changed which steps exist.
   Fixed by duplicating Subprocess_Driver_Agent/Subprocess_Generator_Agent
   verbatim from process.hocon (the same "duplicated not cross-referenced"
   pattern this file's own header comment already explains for
   Design_Agent/Compliance_Agent/Simulation_Agent) and adding a
   coded_tools/process_update/subprocess_tool.py re-export.
"""

import json
import os

import docx
import pytest

from coded_tools.common import paths
from coded_tools.common.process_json import PROCESS_JSON_FILENAME
from coded_tools.process_update.edge_inference_tool import GenerateProcessFlowDiagramCodedTool
from coded_tools.process_update.doc_generation_tool import GenerateProcessDocumentCodedTool
from coded_tools.process_update.subprocess_tool import LoadProcessStepsCodedTool, SaveSubprocessFlowCodedTool


def test_process_update_front_man_has_document_regeneration_tools():
    os.environ.setdefault("AGENT_MANIFEST_FILE", os.path.join(os.getcwd(), "registries", "manifest.hocon"))
    os.environ.setdefault("AGENT_TOOL_PATH", os.path.join(os.getcwd(), "coded_tools"))
    os.environ["PYTHONPATH"] = os.getcwd()

    from neuro_san.internals.graph.persistence.agent_network_restorer import AgentNetworkRestorer

    network = AgentNetworkRestorer().restore(file_reference="registries/process_update.hocon")
    front_man_tools = network.get_config()["tools"][0]["tools"]

    assert "generate_process_flow_diagram" in front_man_tools
    assert "generate_process_document" in front_man_tools


def test_process_update_edge_inference_tool_re_exports_the_real_class():
    # coded_tools/process_update/edge_inference_tool.py is a thin re-export
    # (same pattern as coded_tools/design_update/edge_inference_tool.py),
    # not a fresh declaration -- confirm it resolves to the real class.
    from coded_tools.process.edge_inference_tool import GenerateProcessFlowDiagramCodedTool as Original
    assert GenerateProcessFlowDiagramCodedTool is Original


def test_process_update_doc_generation_tool_re_exports_the_real_class():
    from coded_tools.process.doc_generation_tool import GenerateProcessDocumentCodedTool as Original
    assert GenerateProcessDocumentCodedTool is Original


def test_process_update_coded_tools_actually_regenerate_a_document():
    with open(paths.output_path(PROCESS_JSON_FILENAME), "w", encoding="utf-8") as f:
        json.dump({
            "process_name": "Vendor Onboarding",
            "process_steps": [
                {"step_name": "KYC Check", "responsible_party": "Compliance", "dependencies": []},
            ],
        }, f)

    diagram_result = GenerateProcessFlowDiagramCodedTool().invoke({}, {})
    assert "Diagram successfully generated at" in diagram_result

    doc_result = GenerateProcessDocumentCodedTool().invoke({}, {})
    assert doc_result.startswith("SUCCESS:")
    out_path = os.path.join(paths.OUTPUT_DIR, "Vendor_Onboarding.docx")
    assert os.path.exists(out_path)
    assert docx.Document(out_path).paragraphs  # a real, readable docx was written


def test_process_update_front_man_has_subprocess_driver_agent():
    os.environ.setdefault("AGENT_MANIFEST_FILE", os.path.join(os.getcwd(), "registries", "manifest.hocon"))
    os.environ.setdefault("AGENT_TOOL_PATH", os.path.join(os.getcwd(), "coded_tools"))
    os.environ["PYTHONPATH"] = os.getcwd()

    from neuro_san.internals.graph.persistence.agent_network_restorer import AgentNetworkRestorer

    network = AgentNetworkRestorer().restore(file_reference="registries/process_update.hocon")
    config = network.get_config()
    tool_names = {t["name"] for t in config["tools"]}
    front_man_tools = config["tools"][0]["tools"]

    assert "Subprocess_Driver_Agent" in front_man_tools
    assert {"Subprocess_Driver_Agent", "Subprocess_Generator_Agent",
            "load_process_steps", "save_subprocess_flow"} <= tool_names


def test_process_update_subprocess_tool_re_exports_the_real_classes():
    from coded_tools.process.subprocess_tool import (
        LoadProcessStepsCodedTool as OriginalLoad,
        SaveSubprocessFlowCodedTool as OriginalSave,
    )
    assert LoadProcessStepsCodedTool is OriginalLoad
    assert SaveSubprocessFlowCodedTool is OriginalSave


def test_process_update_subprocess_flow_actually_persists_and_renders():
    with open(paths.output_path(PROCESS_JSON_FILENAME), "w", encoding="utf-8") as f:
        json.dump({
            "process_name": "Vendor Onboarding",
            "process_steps": [
                {"step_name": "KYC Check", "responsible_party": "Compliance", "dependencies": []},
            ],
        }, f)

    steps = LoadProcessStepsCodedTool().invoke({}, {})
    assert steps["process_steps"][0]["step_name"] == "KYC Check"

    flow = {"step_name": "KYC Check", "subprocess_flow": [
        {"substep_name": "Collect Documents", "description": "Gather KYC docs.", "responsible_party": "Compliance"},
    ]}
    save_result = SaveSubprocessFlowCodedTool().invoke({"step_name": "KYC Check", "subprocess_flow": flow}, {})
    assert save_result["status"] == "OK"

    # The real end-to-end point of this fix: a re-expanded subprocess must
    # actually show up in the regenerated document, not just persist to disk.
    doc_result = GenerateProcessDocumentCodedTool().invoke({}, {})
    assert doc_result.startswith("SUCCESS:")
    out_path = os.path.join(paths.OUTPUT_DIR, "Vendor_Onboarding.docx")
    body_text = "\n".join(p.text for p in docx.Document(out_path).paragraphs)
    assert "Collect Documents" in body_text
