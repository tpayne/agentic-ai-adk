from coded_tools.common.directory_extractors import load_directory_context


def test_missing_directory_returns_not_found(tmp_path):
    result = load_directory_context(str(tmp_path / "does-not-exist"))
    assert result["status"] == "NOT_FOUND"


def test_extracts_txt_and_skips_legacy_doc(tmp_path):
    (tmp_path / "notes.txt").write_text("Vendor onboarding requires KYC checks.")
    (tmp_path / "legacy.doc").write_text("ignored -- old binary format")

    result = load_directory_context(str(tmp_path))

    assert result["status"] == "OK"
    assert result["files_processed"] == ["notes.txt"]
    assert "KYC checks" in result["combined_text"]
    assert "=== notes.txt ===" in result["combined_text"]

    skipped = {entry["file"]: entry["reason"] for entry in result["files_skipped"]}
    assert "legacy.doc" in skipped
    assert "re-save as .docx" in skipped["legacy.doc"]


def test_docx_table_text_is_extracted(tmp_path):
    import docx

    doc = docx.Document()
    doc.add_paragraph("Requirements register:")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "REQ-1"
    table.rows[0].cells[1].text = "Support SSO login"
    doc.save(str(tmp_path / "requirements.docx"))

    result = load_directory_context(str(tmp_path))

    assert result["status"] == "OK"
    assert "Support SSO login" in result["combined_text"]


def test_combined_text_is_truncated_past_max_chars(tmp_path):
    (tmp_path / "big.txt").write_text("x" * 1000)

    result = load_directory_context(str(tmp_path), max_chars=50)

    assert result["truncated"] is True
    assert len(result["combined_text"]) < 1000
