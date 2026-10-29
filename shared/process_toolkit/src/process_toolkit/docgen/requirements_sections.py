"""Bespoke renderers for the requirements-tracking document: Introduction
and Purpose (with Assumptions/Risks/Open Questions as subsections),
Stakeholders (reuses content.py's own renderer directly -- no
reimplementation), the Requirements Register (split by type into
Functional/Non-Functional/Goals and Objectives/Constraints tables, or
plain bullets for a legacy old-format summary with no register at all),
Requirement Priorities, Conflicting Requirements, Requirements Needing
Clarification, Inferred Requirements, and Coverage Gaps -- rendered
directly from requirements_summary.json's own fields (see
process_toolkit.requirements_summary.requirements_summary and
process_toolkit.schema.requirements_json for that shape). ISO/IEC/IEEE
29148-style: identified, traceable, prioritized requirements with
acceptance criteria, not a free-form narrative.

Follows design_sections.py's own conventions exactly: a short opening
paragraph, a fixed set of columns, natural-id sort, apply_iso_table_
formatting/_set_cell_bullets, True/False no-op-render return.
"""

import traceback

import docx

from process_toolkit.docgen.structure import apply_iso_table_formatting, _set_cell_bullets, _leading_number
from process_toolkit.docgen.design_sections import _natural_sort_key, _priority_sort_key
from process_toolkit.schema.requirements_json import is_old_format_requirements_summary


def _set_header_row(table, headers) -> None:
    for idx, label in enumerate(headers):
        table.rows[0].cells[idx].text = label


# ============================================================
# 1.0 INTRODUCTION AND PURPOSE
# ============================================================

def _add_requirements_overview_section(
    doc: docx.Document, data: dict, heading: str = "1.0 Introduction and Purpose",
) -> bool:
    """1.0 Introduction and Purpose -- summary prose, source material, and
    industry sector, with Assumptions/Risks/Open Questions as numbered
    subsections -- mirrors content.py's own _add_overview_section shape
    (main narrative + numbered list-field subsections + bolded scalar
    fields). Always renders -- this is the document's own opening
    section, and there is always at least a generic narrative to show."""
    try:
        doc.add_heading(heading, level=1)
        lead = _leading_number(heading, default=1)

        summary = data.get("summary")
        doc.add_paragraph(
            str(summary) if summary else
            "This document records the requirements identified for this effort, their "
            "priorities, known conflicts and open questions, and any gaps identified against "
            "standard requirement categories."
        )

        subsection = 1
        for key, label in (
            ("assumptions", "Assumptions"), ("risks", "Risks"), ("open_questions", "Open Questions"),
        ):
            values = [str(v) for v in (data.get(key) or []) if isinstance(v, str) and v.strip()]
            if not values:
                continue
            doc.add_heading(f"{lead}.{subsection} {label}", level=2)
            subsection += 1
            for v in values:
                doc.add_paragraph(v, style="List Bullet")

        sector = data.get("industry_sector")
        if sector:
            p = doc.add_paragraph()
            p.add_run("Industry / Domain: ").bold = True
            p.add_run(str(sector))

        source_directory = data.get("source_directory")
        source_files = [f for f in (data.get("source_files") or []) if isinstance(f, str) and f.strip()]
        if source_directory or source_files:
            p = doc.add_paragraph()
            p.add_run("Source Material: ").bold = True
            parts = [str(source_directory)] if source_directory else []
            if source_files:
                parts.append(", ".join(source_files))
            p.add_run(" -- ".join(parts))

        return True
    except Exception:
        traceback.print_exc()
        return False


# ============================================================
# 3.0 REQUIREMENTS REGISTER
# ============================================================

_TYPE_SECTIONS = (
    ("functional", "Functional Requirements"),
    ("non_functional", "Non-Functional Requirements"),
    ("goal", "Goals and Objectives"),
    ("constraint", "Constraints"),
)

_OLD_FORMAT_SECTIONS = (
    ("functional_requirements", "Functional Requirements"),
    ("non_functional_requirements", "Non-Functional Requirements"),
    ("goals_or_objectives", "Goals and Objectives"),
    ("constraints", "Constraints"),
)


def _add_requirements_register_table(doc: docx.Document, items) -> None:
    items = sorted(items, key=lambda r: _natural_sort_key(r.get("id")))

    headers = ["ID", "Description", "Source", "Acceptance Criteria", "Status"]
    table = doc.add_table(rows=1, cols=len(headers))
    _set_header_row(table, headers)

    for item in items:
        row = table.add_row().cells
        row[0].text = str(item.get("id") or "N/A")
        _set_cell_bullets(row[1], item.get("description"))
        row[2].text = str(item.get("source") or "")
        _set_cell_bullets(row[3], item.get("acceptance_criteria"))
        row[4].text = str(item.get("status") or "").replace("_", " ").title()

    apply_iso_table_formatting(table, doc)
    doc.add_paragraph()


def _add_old_format_requirements_bullets(doc: docx.Document, data: dict, lead: int) -> bool:
    """Legacy flat-string-array shape -- no ids/source/acceptance criteria
    to tabulate, so each category renders as plain bullets instead."""
    rendered_any = False
    subsection = 1
    for key, label in _OLD_FORMAT_SECTIONS:
        values = [str(v) for v in (data.get(key) or []) if isinstance(v, str) and v.strip()]
        if not values:
            continue
        doc.add_heading(f"{lead}.{subsection} {label}", level=2)
        subsection += 1
        for v in values:
            doc.add_paragraph(v, style="List Bullet")
        rendered_any = True
    return rendered_any


def _add_requirements_register_section(
    doc: docx.Document, data: dict, heading: str = "3.0 Requirements Register",
) -> bool:
    """3.0 Requirements Register -- ISO/IEC/IEEE 29148 style. Splits the
    register by type into Functional/Non-Functional/Goals and Objectives/
    Constraints tables (ID/Description/Source/Acceptance Criteria/
    Status), sorted by id. Falls back to plain bulleted lists for a
    legacy old-format summary with no register at all (see
    is_old_format_requirements_summary)."""
    try:
        lead = _leading_number(heading, default=3)

        if is_old_format_requirements_summary(data):
            doc.add_heading(heading, level=1)
            doc.add_paragraph(
                "This requirements set predates per-requirement identifiers and traceability "
                "fields; each category below is listed without an id, source, or acceptance "
                "criteria."
            )
            return _add_old_format_requirements_bullets(doc, data, lead)

        register = [r for r in (data.get("requirements_register") or []) if isinstance(r, dict)]
        if not register:
            return False

        by_type: dict = {}
        for item in register:
            by_type.setdefault(item.get("type"), []).append(item)

        doc.add_heading(heading, level=1)
        doc.add_paragraph(
            "This section identifies the functional, non-functional, goal, and constraint "
            "requirements for this effort, following ISO/IEC/IEEE 29148 requirements "
            "engineering practice. Each table is ordered by identifier."
        )

        subsection = 1
        for req_type, label in _TYPE_SECTIONS:
            items = by_type.get(req_type) or []
            if not items:
                continue
            doc.add_heading(f"{lead}.{subsection} {label}", level=2)
            subsection += 1
            _add_requirements_register_table(doc, items)

        return True
    except Exception:
        traceback.print_exc()
        return False


# ============================================================
# 4.0 REQUIREMENT PRIORITIES
# ============================================================

def _add_requirement_priorities_section(
    doc: docx.Document, priorities, heading: str = "4.0 Requirement Priorities",
) -> bool:
    """4.0 Requirement Priorities -- one table (ID/Requirement/Priority/
    Rationale), sorted by priority then id."""
    try:
        items = [p for p in (priorities or []) if isinstance(p, dict)]
        if not items:
            return False

        doc.add_heading(heading, level=1)
        doc.add_paragraph(
            "This section records the relative priority assigned to each requirement, and the "
            "rationale behind it."
        )

        items = sorted(
            items,
            key=lambda p: (_priority_sort_key(p.get("priority")), _natural_sort_key(p.get("requirement_id"))),
        )

        headers = ["ID", "Requirement", "Priority", "Rationale"]
        table = doc.add_table(rows=1, cols=len(headers))
        _set_header_row(table, headers)
        for p in items:
            row = table.add_row().cells
            row[0].text = str(p.get("requirement_id") or "")
            _set_cell_bullets(row[1], p.get("requirement_text"))
            row[2].text = str(p.get("priority") or "").title()
            _set_cell_bullets(row[3], p.get("rationale"))
        apply_iso_table_formatting(table, doc)
        doc.add_paragraph()
        return True
    except Exception:
        traceback.print_exc()
        return False


# ============================================================
# 5.0 CONFLICTING REQUIREMENTS
# ============================================================

def _add_conflicting_requirements_section(
    doc: docx.Document, conflicts, heading: str = "5.0 Conflicting Requirements",
) -> bool:
    """5.0 Conflicting Requirements -- one table (Requirement IDs/
    Requirements/Description/Recommendation)."""
    try:
        items = [c for c in (conflicts or []) if isinstance(c, dict)]
        if not items:
            return False

        doc.add_heading(heading, level=1)
        doc.add_paragraph(
            "This section records requirements identified as being in conflict with one "
            "another, and a recommended resolution."
        )

        headers = ["Requirement IDs", "Requirements", "Description", "Recommendation"]
        table = doc.add_table(rows=1, cols=len(headers))
        _set_header_row(table, headers)
        for c in items:
            row = table.add_row().cells
            ids = [str(i) for i in (c.get("requirement_ids") or [])]
            row[0].text = ", ".join(ids)
            _set_cell_bullets(row[1], c.get("requirements_text"))
            _set_cell_bullets(row[2], c.get("description"))
            _set_cell_bullets(row[3], c.get("recommendation"))
        apply_iso_table_formatting(table, doc)
        doc.add_paragraph()
        return True
    except Exception:
        traceback.print_exc()
        return False


# ============================================================
# 6.0 REQUIREMENTS NEEDING CLARIFICATION
# ============================================================

def _add_requirements_needing_clarification_section(
    doc: docx.Document, items, heading: str = "6.0 Requirements Needing Clarification",
) -> bool:
    """6.0 Requirements Needing Clarification -- one table (ID/
    Requirement/Question), sorted by id."""
    try:
        items = [i for i in (items or []) if isinstance(i, dict)]
        if not items:
            return False

        doc.add_heading(heading, level=1)
        doc.add_paragraph(
            "This section records requirements that are too ambiguous to act on as stated, "
            "and the clarifying question that needs to be answered."
        )

        items = sorted(items, key=lambda i: _natural_sort_key(i.get("requirement_id")))

        headers = ["ID", "Requirement", "Question"]
        table = doc.add_table(rows=1, cols=len(headers))
        _set_header_row(table, headers)
        for i in items:
            row = table.add_row().cells
            row[0].text = str(i.get("requirement_id") or "")
            _set_cell_bullets(row[1], i.get("requirement_text"))
            _set_cell_bullets(row[2], i.get("question"))
        apply_iso_table_formatting(table, doc)
        doc.add_paragraph()
        return True
    except Exception:
        traceback.print_exc()
        return False


# ============================================================
# 7.0 INFERRED REQUIREMENTS
# ============================================================

def _add_inferred_requirements_section(
    doc: docx.Document, items, heading: str = "7.0 Inferred Requirements",
) -> bool:
    """7.0 Inferred Requirements -- requirements the analysis stage judged
    were likely needed but were NOT actually stated in the source
    material. Opens with a mandatory disclaimer (mirrors Requirements_
    Consultant_Agent's own labeling rule) that these are analyst-inferred,
    never sourced, and must not be mistaken for confirmed requirements."""
    try:
        items = [i for i in (items or []) if isinstance(i, dict)]
        if not items:
            return False

        doc.add_heading(heading, level=1)
        doc.add_paragraph(
            "The following requirements were NOT stated in the source material -- they were "
            "inferred as likely necessary based on the stated requirements and context. They "
            "are not confirmed and must be validated with stakeholders before being treated as "
            "committed requirements."
        )

        items = sorted(items, key=lambda i: _natural_sort_key(i.get("id")))

        headers = ["ID", "Requirement", "Rationale", "Category", "Derived From", "Acceptance Criteria"]
        table = doc.add_table(rows=1, cols=len(headers))
        _set_header_row(table, headers)
        for i in items:
            row = table.add_row().cells
            row[0].text = str(i.get("id") or "")
            _set_cell_bullets(row[1], i.get("requirement"))
            _set_cell_bullets(row[2], i.get("rationale"))
            row[3].text = str(i.get("category") or "").replace("_", " ").title()
            derived_from = [str(d) for d in (i.get("derived_from") or [])]
            row[4].text = ", ".join(derived_from)
            _set_cell_bullets(row[5], i.get("acceptance_criteria"))
        apply_iso_table_formatting(table, doc)
        doc.add_paragraph()
        return True
    except Exception:
        traceback.print_exc()
        return False


# ============================================================
# 8.0 COVERAGE GAPS
# ============================================================

def _add_coverage_gaps_section(
    doc: docx.Document, gaps, heading: str = "8.0 Coverage Gaps",
) -> bool:
    """8.0 Coverage Gaps -- standard requirement categories the source
    material does not appear to address at all, one table
    (Category/Description) sorted alphabetically by category."""
    try:
        items = [g for g in (gaps or []) if isinstance(g, dict)]
        if not items:
            return False

        doc.add_heading(heading, level=1)
        doc.add_paragraph(
            "This section lists standard requirement categories that the source material does "
            "not appear to address at all."
        )

        items = sorted(items, key=lambda g: str(g.get("category") or "").lower())

        headers = ["Category", "Description"]
        table = doc.add_table(rows=1, cols=len(headers))
        _set_header_row(table, headers)
        for g in items:
            row = table.add_row().cells
            row[0].text = str(g.get("category") or "")
            _set_cell_bullets(row[1], g.get("description"))
        apply_iso_table_formatting(table, doc)
        doc.add_paragraph()
        return True
    except Exception:
        traceback.print_exc()
        return False
