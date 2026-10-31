"""Bespoke renderers for the proposal-response document -- a full
enterprise-RTP table of contents: Executive Summary (with Understanding/
Solution Summary/Differentiators/Expected-Outcomes subsections), Proposed
Solution, Solution Architecture (embeds a generated component diagram),
Scope of Work & Boundaries (Services in Scope/computed Functional
Coverage/Environments/Deliverables/Out-of-Scope/Technical Constraints),
the Compliance Matrix (the central traceability table, one row per
requirement id from the SOURCE requirements_summary.json that
Proposal_Response_Agent drafted this response against, each with an
Evidence/Location pointer into this same document), Delivery &
Engagement Approach (domain expertise/accelerators/core approach/
platform approach/quality validation/automation/environment setup/
tooling stack/SLAs/mobilisation roadmap), Governance Model (framework/
RACI matrix/escalation paths/meeting cadence), Indicative Profiles &
Team Structure (team structure/key personnel/skills & certifications),
Relevant Success Stories & References (case studies/measurable
outcomes), Qualifications, Key Assumptions & Dependencies (assumptions/
client dependencies/risks & mitigations), Commercial Proposal & TCO
(pricing approach plus itemized pricing table/ongoing costs/rate cards),
Open Questions for Clarification, and Unaddressed Items -- rendered
directly from proposal_response.json's own fields (see
process_toolkit.proposal_response.proposal_response and
process_toolkit.schema.proposal_response_json for that shape).

Follows requirements_sections.py's own conventions exactly: a short
opening paragraph, a fixed set of columns, natural-id sort,
apply_iso_table_formatting/_set_cell_bullets, True/False no-op-render
return. Container sections (e.g. Scope of Work, Delivery & Engagement
Approach) are a single Heading-1 holding several Heading-2 subsections,
each independently no-op when its own backing field is absent -- mirrors
design_sections.py's own nested-heading convention. A container section
itself is a no-op (no heading at all) only when EVERY one of its own
subsections is also empty -- Executive Summary is the one exception,
since its own opening paragraph always has a generic fallback.
"""

import traceback

import docx
from docx.shared import Inches

from process_toolkit.docgen.structure import apply_iso_table_formatting, _set_cell_bullets
from process_toolkit.docgen.design_sections import _natural_sort_key
from process_toolkit.docgen.proposal_architecture_diagram import generate_proposal_architecture_diagram


def _set_header_row(table, headers) -> None:
    for idx, label in enumerate(headers):
        table.rows[0].cells[idx].text = label


def _add_bullet_list_subsection(doc: docx.Document, values, heading: str, level: int = 2) -> bool:
    """Shared helper for every plain string-list subsection introduced in
    this pass (services_in_scope, environments_covered, out_of_scope,
    technical_constraints, domain_accelerators, tooling_stack,
    skills_and_certifications, measurable_outcomes, client_dependencies,
    escalation_path) -- same no-op/bullet-list shape as the pre-existing
    differentiators/assumptions/open_questions renderers, just
    parameterized so it isn't copy-pasted ten more times."""
    try:
        items = [str(v) for v in (values or []) if isinstance(v, str) and v.strip()]
        if not items:
            return False
        doc.add_heading(heading, level=level)
        for v in items:
            doc.add_paragraph(v, style="List Bullet")
        return True
    except Exception:
        traceback.print_exc()
        return False


def _add_prose_subsection(doc: docx.Document, text, heading: str, level: int = 2) -> bool:
    """Shared helper for every single-paragraph prose subsection
    introduced in this pass (domain_expertise, core_delivery_approach,
    platform_approach, quality_validation_approach, automation_approach,
    environment_setup, governance_framework, ongoing_costs_approach,
    rate_card_and_change_control)."""
    try:
        if not text:
            return False
        doc.add_heading(heading, level=level)
        doc.add_paragraph(str(text))
        return True
    except Exception:
        traceback.print_exc()
        return False


def _add_table_subsection(
    doc: docx.Document, items, headers, row_keys, heading: str, level: int = 2, plain_columns=None,
) -> bool:
    """Shared helper for every simple same-shape table subsection
    introduced in this pass (slas_and_kpis, raci_matrix, meeting_cadence,
    team_structure, key_personnel). `plain_columns` is the set of column
    indices that get a direct `.text =` assignment (short, identifier-
    like values -- a role name, a metric label) rather than
    _set_cell_bullets (longer prose that may wrap as multiple bullets)
    -- mirrors the pre-existing pattern of e.g. the compliance matrix's
    own ID/Status columns being plain while Requirement/Response are
    bulleted."""
    try:
        rows = [r for r in (items or []) if isinstance(r, dict)]
        if not rows:
            return False
        plain_columns = plain_columns or set()
        doc.add_heading(heading, level=level)
        table = doc.add_table(rows=1, cols=len(headers))
        _set_header_row(table, headers)
        for r in rows:
            cells = table.add_row().cells
            for idx, key in enumerate(row_keys):
                if idx in plain_columns:
                    cells[idx].text = str(r.get(key) or "")
                else:
                    _set_cell_bullets(cells[idx], r.get(key))
        apply_iso_table_formatting(table, doc)
        doc.add_paragraph()
        return True
    except Exception:
        traceback.print_exc()
        return False


# ============================================================
# 1.0 EXECUTIVE SUMMARY
# ============================================================

def _add_understanding_section(
    doc: docx.Document, data: dict, heading: str = "1.1 Understanding of Client Objectives",
) -> bool:
    try:
        understanding = data.get("our_understanding")
        if not understanding:
            return False
        doc.add_heading(heading, level=2)
        doc.add_paragraph(str(understanding))
        return True
    except Exception:
        traceback.print_exc()
        return False


def _add_solution_summary_section(
    doc: docx.Document, data: dict, heading: str = "1.2 Summary of Proposed Solution",
) -> bool:
    return _add_prose_subsection(doc, data.get("solution_summary"), heading)


def _add_differentiators_section(
    doc: docx.Document, differentiators, heading: str = "1.3 Differentiators & Value Proposition",
) -> bool:
    return _add_bullet_list_subsection(doc, differentiators, heading)


def _add_expected_outcomes_section(
    doc: docx.Document, outcomes, heading: str = "1.4 Expected Outcomes & ROI",
) -> bool:
    return _add_bullet_list_subsection(doc, outcomes, heading)


def _add_executive_summary_section(
    doc: docx.Document, data: dict, heading: str = "1.0 Executive Summary",
) -> bool:
    """1.0 Executive Summary -- always renders, same as requirements_
    sections.py's own Introduction and Purpose section: this is the
    document's own opening section, and there is always at least a
    generic narrative to show. Its four subsections are each optional,
    no-op when their own backing field is absent."""
    try:
        doc.add_heading(heading, level=1)
        summary = data.get("summary")
        doc.add_paragraph(
            str(summary) if summary else
            "This document sets out our response to the requirements identified for this "
            "proposal, including our understanding of those requirements, our proposed "
            "solution, and a point-by-point compliance matrix."
        )
        _add_understanding_section(doc, data)
        _add_solution_summary_section(doc, data)
        _add_differentiators_section(doc, data.get("differentiators"))
        _add_expected_outcomes_section(doc, data.get("expected_outcomes"))
        return True
    except Exception:
        traceback.print_exc()
        return False


# ============================================================
# 2.0 PROPOSED SOLUTION
# ============================================================

def _add_proposed_solution_section(
    doc: docx.Document, data: dict, heading: str = "2.0 Proposed Solution",
) -> bool:
    try:
        solution = data.get("proposed_solution")
        if not solution:
            return False
        doc.add_heading(heading, level=1)
        doc.add_paragraph(str(solution))
        return True
    except Exception:
        traceback.print_exc()
        return False


# ============================================================
# 3.0 SOLUTION ARCHITECTURE
# ============================================================

def _add_solution_architecture_section(
    doc: docx.Document, components, process_name: str, heading: str = "3.0 Solution Architecture",
) -> bool:
    """3.0 Solution Architecture -- renders a generated component diagram
    (see proposal_architecture_diagram.py) plus a Component/Description
    table underneath for accessibility/non-image contexts. No-op (no
    fabricated section) when "architecture_components" has no usable
    entries -- same convention as every other renderer here."""
    try:
        items = [c for c in (components or []) if isinstance(c, dict) and c.get("name")]
        if not items:
            return False

        diagram_path = generate_proposal_architecture_diagram(items, process_name)

        doc.add_heading(heading, level=1)
        doc.add_paragraph(
            "The following diagram provides a high-level visualization of the proposed solution's "
            "architecture."
        )
        doc.add_picture(diagram_path, width=Inches(6))
        doc.add_paragraph()

        headers = ["Component", "Description"]
        table = doc.add_table(rows=1, cols=len(headers))
        _set_header_row(table, headers)
        for item in items:
            row = table.add_row().cells
            row[0].text = str(item.get("name") or "")
            _set_cell_bullets(row[1], item.get("description"))
        apply_iso_table_formatting(table, doc)
        doc.add_paragraph()
        return True
    except Exception:
        traceback.print_exc()
        return False


# ============================================================
# 4.0 SCOPE OF WORK & BOUNDARIES
# ============================================================

def _add_services_in_scope_section(
    doc: docx.Document, services, heading: str = "4.1 Services in Scope",
) -> bool:
    return _add_bullet_list_subsection(doc, services, heading)


def _add_functional_coverage_section(
    doc: docx.Document, matrix, heading: str = "4.2 Functional Coverage",
) -> bool:
    """4.2 Functional Coverage -- a COMPUTED rollup, not an agent-drafted
    field: counts compliance_matrix's own response_status values
    directly, so this carries zero fabrication risk and needs no new
    schema field at all."""
    try:
        items = [m for m in (matrix or []) if isinstance(m, dict)]
        if not items:
            return False
        total = len(items)
        counts: dict = {}
        for m in items:
            status = str(m.get("response_status") or "Unspecified")
            counts[status] = counts.get(status, 0) + 1
        status_summary = ", ".join(f"{count} {status}" for status, count in counts.items())
        doc.add_heading(heading, level=2)
        doc.add_paragraph(
            f"This proposal provides a response to all {total} requirements identified in the "
            f"requirements register: {status_summary}."
        )
        return True
    except Exception:
        traceback.print_exc()
        return False


def _add_environments_covered_section(
    doc: docx.Document, environments, heading: str = "4.3 Environments & Platforms Covered",
) -> bool:
    return _add_bullet_list_subsection(doc, environments, heading)


def _add_deliverables_section(
    doc: docx.Document, deliverables, heading: str = "4.4 Key Deliverables & Artifacts",
) -> bool:
    return _add_bullet_list_subsection(doc, deliverables, heading)


def _add_out_of_scope_section(
    doc: docx.Document, items, heading: str = "4.5 Out-of-Scope Activities",
) -> bool:
    return _add_bullet_list_subsection(doc, items, heading)


def _add_technical_constraints_section(
    doc: docx.Document, items, heading: str = "4.6 Key Technical Constraints",
) -> bool:
    return _add_bullet_list_subsection(doc, items, heading)


def _add_scope_of_work_section(
    doc: docx.Document, data: dict, heading: str = "4.0 Scope of Work & Boundaries",
) -> bool:
    """4.0 Scope of Work & Boundaries -- a container section: the parent
    heading is only added if at least one of its own subsections' raw
    backing field is truthy (checked up front, same convention every
    other optional section here already uses), then each subsection
    renders itself (and independently no-ops if its own field turns out
    empty after filtering, e.g. a list of blank strings)."""
    try:
        if not any([
            data.get("services_in_scope"), data.get("compliance_matrix"),
            data.get("environments_covered"), data.get("deliverables"),
            data.get("out_of_scope"), data.get("technical_constraints"),
        ]):
            return False

        doc.add_heading(heading, level=1)
        _add_services_in_scope_section(doc, data.get("services_in_scope"))
        _add_functional_coverage_section(doc, data.get("compliance_matrix"))
        _add_environments_covered_section(doc, data.get("environments_covered"))
        _add_deliverables_section(doc, data.get("deliverables"))
        _add_out_of_scope_section(doc, data.get("out_of_scope"))
        _add_technical_constraints_section(doc, data.get("technical_constraints"))
        return True
    except Exception:
        traceback.print_exc()
        return False


# ============================================================
# 5.0 COMPLIANCE MATRIX
# ============================================================

def _add_compliance_matrix_section(
    doc: docx.Document, matrix, heading: str = "5.0 Compliance Matrix",
) -> bool:
    """5.0 Compliance Matrix -- the central artifact of this document: one
    row per requirement id carried over from the SOURCE
    requirements_summary.json's own requirements_register, sorted by
    that same id, so a reader can trace every requirement to a response."""
    try:
        items = [m for m in (matrix or []) if isinstance(m, dict)]
        if not items:
            return False

        doc.add_heading(heading, level=1)
        doc.add_paragraph(
            "This section addresses each identified requirement in turn, stating our "
            "compliance status and response."
        )

        items = sorted(items, key=lambda m: _natural_sort_key(m.get("requirement_id")))

        headers = ["ID", "Requirement", "Status", "Response", "Evidence / Location"]
        table = doc.add_table(rows=1, cols=len(headers))
        _set_header_row(table, headers)
        for m in items:
            row = table.add_row().cells
            row[0].text = str(m.get("requirement_id") or "")
            _set_cell_bullets(row[1], m.get("requirement_text"))
            row[2].text = str(m.get("response_status") or "")
            _set_cell_bullets(row[3], m.get("response"))
            _set_cell_bullets(row[4], m.get("evidence_ref"))
        apply_iso_table_formatting(table, doc)
        doc.add_paragraph()
        return True
    except Exception:
        traceback.print_exc()
        return False


# ============================================================
# 6.0 DELIVERY & ENGAGEMENT APPROACH
# ============================================================

def _add_domain_expertise_section(
    doc: docx.Document, data: dict, heading: str = "6.1 Industry / Domain Expertise",
) -> bool:
    return _add_prose_subsection(doc, data.get("domain_expertise"), heading)


def _add_domain_accelerators_section(
    doc: docx.Document, items, heading: str = "6.2 Domain-Specific Accelerators",
) -> bool:
    return _add_bullet_list_subsection(doc, items, heading)


def _add_core_delivery_approach_section(
    doc: docx.Document, data: dict, heading: str = "6.3 Core Delivery Approach",
) -> bool:
    return _add_prose_subsection(doc, data.get("core_delivery_approach"), heading)


def _add_platform_approach_section(
    doc: docx.Document, data: dict, heading: str = "6.4 Platform / Application Approach",
) -> bool:
    return _add_prose_subsection(doc, data.get("platform_approach"), heading)


def _add_quality_validation_approach_section(
    doc: docx.Document, data: dict, heading: str = "6.5 Data & Quality Validation Approach",
) -> bool:
    return _add_prose_subsection(doc, data.get("quality_validation_approach"), heading)


def _add_automation_approach_section(
    doc: docx.Document, data: dict, heading: str = "6.6 Automation & Tooling Approach",
) -> bool:
    return _add_prose_subsection(doc, data.get("automation_approach"), heading)


def _add_environment_setup_section(
    doc: docx.Document, data: dict, heading: str = "6.7 Environment Setup & Hosting",
) -> bool:
    return _add_prose_subsection(doc, data.get("environment_setup"), heading)


def _add_tooling_stack_section(
    doc: docx.Document, items, heading: str = "6.8 Recommended Tooling Stack",
) -> bool:
    return _add_bullet_list_subsection(doc, items, heading)


def _add_slas_and_kpis_section(
    doc: docx.Document, items, heading: str = "6.9 SLAs & Quality KPIs",
) -> bool:
    return _add_table_subsection(
        doc, items, ["Metric", "Target"], ["metric", "target"], heading, plain_columns={0},
    )


def _add_implementation_plan_section(
    doc: docx.Document, plan, heading: str = "6.10 Mobilisation & Phased Roadmap",
) -> bool:
    """6.10 Mobilisation & Phased Roadmap -- a phased delivery timeline.
    Rows are left in the agent's own given order (not sorted) since
    phases are inherently sequential, unlike the compliance matrix's
    requirement ids."""
    try:
        items = [p for p in (plan or []) if isinstance(p, dict)]
        if not items:
            return False

        doc.add_heading(heading, level=2)
        headers = ["Phase", "Timing", "Deliverable", "Acceptance Criteria"]
        table = doc.add_table(rows=1, cols=len(headers))
        _set_header_row(table, headers)
        for p in items:
            row = table.add_row().cells
            row[0].text = str(p.get("phase") or "")
            row[1].text = str(p.get("timing") or "")
            _set_cell_bullets(row[2], p.get("deliverable"))
            _set_cell_bullets(row[3], p.get("acceptance_criteria"))
        apply_iso_table_formatting(table, doc)
        doc.add_paragraph()
        return True
    except Exception:
        traceback.print_exc()
        return False


def _add_delivery_approach_section(
    doc: docx.Document, data: dict, heading: str = "6.0 Delivery & Engagement Approach",
) -> bool:
    try:
        if not any([
            data.get("domain_expertise"), data.get("domain_accelerators"),
            data.get("core_delivery_approach"), data.get("platform_approach"),
            data.get("quality_validation_approach"), data.get("automation_approach"),
            data.get("environment_setup"), data.get("tooling_stack"),
            data.get("slas_and_kpis"), data.get("implementation_plan"),
        ]):
            return False

        doc.add_heading(heading, level=1)
        _add_domain_expertise_section(doc, data)
        _add_domain_accelerators_section(doc, data.get("domain_accelerators"))
        _add_core_delivery_approach_section(doc, data)
        _add_platform_approach_section(doc, data)
        _add_quality_validation_approach_section(doc, data)
        _add_automation_approach_section(doc, data)
        _add_environment_setup_section(doc, data)
        _add_tooling_stack_section(doc, data.get("tooling_stack"))
        _add_slas_and_kpis_section(doc, data.get("slas_and_kpis"))
        _add_implementation_plan_section(doc, data.get("implementation_plan"))
        return True
    except Exception:
        traceback.print_exc()
        return False


# ============================================================
# 7.0 GOVERNANCE MODEL, ROLES & RESPONSIBILITIES
# ============================================================

def _add_governance_framework_section(
    doc: docx.Document, data: dict, heading: str = "7.1 Governance Framework",
) -> bool:
    return _add_prose_subsection(doc, data.get("governance_framework"), heading)


def _add_raci_matrix_section(
    doc: docx.Document, items, heading: str = "7.2 RACI Matrix",
) -> bool:
    return _add_table_subsection(
        doc, items,
        ["Activity", "Responsible", "Accountable", "Consulted", "Informed"],
        ["activity", "responsible", "accountable", "consulted", "informed"],
        heading, plain_columns={1, 2, 3, 4},
    )


def _add_escalation_path_section(
    doc: docx.Document, items, heading: str = "7.3 Escalation Paths & Decision-Making",
) -> bool:
    return _add_bullet_list_subsection(doc, items, heading)


def _add_meeting_cadence_section(
    doc: docx.Document, items, heading: str = "7.4 Meeting Cadence & Stakeholder Engagement",
) -> bool:
    return _add_table_subsection(
        doc, items, ["Meeting", "Frequency", "Attendees"], ["meeting", "frequency", "attendees"], heading,
        plain_columns={0, 1},
    )


def _add_governance_section(
    doc: docx.Document, data: dict, heading: str = "7.0 Governance Model, Roles & Responsibilities",
) -> bool:
    try:
        if not any([
            data.get("governance_framework"), data.get("raci_matrix"),
            data.get("escalation_path"), data.get("meeting_cadence"),
        ]):
            return False

        doc.add_heading(heading, level=1)
        _add_governance_framework_section(doc, data)
        _add_raci_matrix_section(doc, data.get("raci_matrix"))
        _add_escalation_path_section(doc, data.get("escalation_path"))
        _add_meeting_cadence_section(doc, data.get("meeting_cadence"))
        return True
    except Exception:
        traceback.print_exc()
        return False


# ============================================================
# 8.0 INDICATIVE PROFILES & TEAM STRUCTURE
# ============================================================

def _add_team_structure_section(
    doc: docx.Document, items, heading: str = "8.1 Team Structure & Allocation",
) -> bool:
    return _add_table_subsection(
        doc, items, ["Role", "Allocation"], ["role", "allocation"], heading, plain_columns={0, 1},
    )


def _add_key_personnel_section(
    doc: docx.Document, items, heading: str = "8.2 Key Personnel Profiles",
) -> bool:
    """8.2 Key Personnel Profiles -- ROLE-level descriptions only (e.g.
    "Delivery Lead", "QE Automation Engineer"), never a fabricated named
    individual or CV -- enforced at the drafting-instruction level, not
    here; this renderer just displays whatever role/responsibilities
    pairs are present."""
    return _add_table_subsection(
        doc, items, ["Role", "Responsibilities"], ["role", "responsibilities"], heading, plain_columns={0},
    )


def _add_skills_and_certifications_section(
    doc: docx.Document, items, heading: str = "8.3 Skills Matrix & Certifications",
) -> bool:
    return _add_bullet_list_subsection(doc, items, heading)


def _add_team_profiles_section(
    doc: docx.Document, data: dict, heading: str = "8.0 Indicative Profiles & Team Structure",
) -> bool:
    try:
        if not any([
            data.get("team_structure"), data.get("key_personnel"), data.get("skills_and_certifications"),
        ]):
            return False

        doc.add_heading(heading, level=1)
        _add_team_structure_section(doc, data.get("team_structure"))
        _add_key_personnel_section(doc, data.get("key_personnel"))
        _add_skills_and_certifications_section(doc, data.get("skills_and_certifications"))
        return True
    except Exception:
        traceback.print_exc()
        return False


# ============================================================
# 9.0 RELEVANT SUCCESS STORIES & REFERENCES
# ============================================================

def _add_case_studies_section(
    doc: docx.Document, items, heading: str = "9.1 Case Studies",
) -> bool:
    """9.1 Case Studies -- the highest fabrication-risk field in this
    document: each case study's title renders as a bold paragraph (not a
    numbered heading, so no renumbering complexity), followed by its own
    Context/Challenge/Work Completed/Outcome fields. Populating a case
    study the user's own input doesn't actually describe is a drafting-
    instruction concern, not something this renderer can police -- it
    renders whatever is given, same as every other renderer here."""
    try:
        rows = [c for c in (items or []) if isinstance(c, dict) and c.get("title")]
        if not rows:
            return False
        doc.add_heading(heading, level=2)
        labels = (
            ("context", "Context"),
            ("challenge", "Challenge"),
            ("work_completed", "Work Completed"),
            ("outcome", "Outcome"),
        )
        for case in rows:
            title_p = doc.add_paragraph()
            title_p.add_run(str(case.get("title"))).bold = True
            for key, label in labels:
                value = case.get(key)
                if value:
                    doc.add_paragraph(f"{label}: {value}")
        return True
    except Exception:
        traceback.print_exc()
        return False


def _add_measurable_outcomes_section(
    doc: docx.Document, items, heading: str = "9.2 Measurable Outcomes & Testimonials",
) -> bool:
    return _add_bullet_list_subsection(doc, items, heading)


def _add_success_stories_section(
    doc: docx.Document, data: dict, heading: str = "9.0 Relevant Success Stories & References",
) -> bool:
    try:
        if not any([data.get("case_studies"), data.get("measurable_outcomes")]):
            return False

        doc.add_heading(heading, level=1)
        _add_case_studies_section(doc, data.get("case_studies"))
        _add_measurable_outcomes_section(doc, data.get("measurable_outcomes"))
        return True
    except Exception:
        traceback.print_exc()
        return False


# ============================================================
# 10.0 QUALIFICATIONS
# ============================================================

def _add_qualifications_section(
    doc: docx.Document, data: dict, heading: str = "10.0 Qualifications",
) -> bool:
    try:
        qualifications = data.get("qualifications")
        if not qualifications:
            return False
        doc.add_heading(heading, level=1)
        doc.add_paragraph(str(qualifications))
        return True
    except Exception:
        traceback.print_exc()
        return False


# ============================================================
# 11.0 KEY ASSUMPTIONS & DEPENDENCIES
# ============================================================

def _add_assumptions_section(
    doc: docx.Document, assumptions, heading: str = "11.1 Project Assumptions",
) -> bool:
    return _add_bullet_list_subsection(doc, assumptions, heading)


def _add_client_dependencies_section(
    doc: docx.Document, items, heading: str = "11.2 Client Responsibilities & Dependencies",
) -> bool:
    return _add_bullet_list_subsection(doc, items, heading)


def _add_risks_and_mitigations_section(
    doc: docx.Document, risks, heading: str = "11.3 Risk Management & Mitigation Matrix",
) -> bool:
    try:
        items = [r for r in (risks or []) if isinstance(r, dict)]
        if not items:
            return False

        doc.add_heading(heading, level=2)
        headers = ["Risk", "Mitigation"]
        table = doc.add_table(rows=1, cols=len(headers))
        _set_header_row(table, headers)
        for r in items:
            row = table.add_row().cells
            _set_cell_bullets(row[0], r.get("risk"))
            _set_cell_bullets(row[1], r.get("mitigation"))
        apply_iso_table_formatting(table, doc)
        doc.add_paragraph()
        return True
    except Exception:
        traceback.print_exc()
        return False


def _add_assumptions_dependencies_section(
    doc: docx.Document, data: dict, heading: str = "11.0 Key Assumptions & Dependencies",
) -> bool:
    try:
        if not any([
            data.get("assumptions"), data.get("client_dependencies"), data.get("risks_and_mitigations"),
        ]):
            return False

        doc.add_heading(heading, level=1)
        _add_assumptions_section(doc, data.get("assumptions"))
        _add_client_dependencies_section(doc, data.get("client_dependencies"))
        _add_risks_and_mitigations_section(doc, data.get("risks_and_mitigations"))
        return True
    except Exception:
        traceback.print_exc()
        return False


# ============================================================
# 12.0 COMMERCIAL PROPOSAL & TOTAL COST OF OWNERSHIP (TCO)
# ============================================================

def _add_pricing_approach_section(
    doc: docx.Document, data: dict, heading: str = "12.1 Fixed-Fee Implementation Schedule",
) -> bool:
    """12.1 Fixed-Fee Implementation Schedule -- the pricing-model prose
    plus, when present, an itemized pricing_table breakdown underneath.
    Renders if EITHER is present; a line item with no actual "price" the
    user ever gave renders as "TBD" rather than being fabricated or
    dropped."""
    try:
        pricing = data.get("pricing_approach")
        table_items = [p for p in (data.get("pricing_table") or []) if isinstance(p, dict)]
        if not pricing and not table_items:
            return False

        doc.add_heading(heading, level=2)
        if pricing:
            doc.add_paragraph(str(pricing))

        if table_items:
            headers = ["Item", "Basis", "Price"]
            table = doc.add_table(rows=1, cols=len(headers))
            _set_header_row(table, headers)
            for item in table_items:
                row = table.add_row().cells
                row[0].text = str(item.get("item") or "")
                _set_cell_bullets(row[1], item.get("basis"))
                row[2].text = str(item.get("price") or "TBD")
            apply_iso_table_formatting(table, doc)
            doc.add_paragraph()
        return True
    except Exception:
        traceback.print_exc()
        return False


def _add_ongoing_costs_approach_section(
    doc: docx.Document, data: dict, heading: str = "12.2 Ongoing Managed Support & Operational Expenses",
) -> bool:
    return _add_prose_subsection(doc, data.get("ongoing_costs_approach"), heading)


def _add_rate_card_and_change_control_section(
    doc: docx.Document, data: dict, heading: str = "12.3 Rate Cards & Change Control Mechanics",
) -> bool:
    return _add_prose_subsection(doc, data.get("rate_card_and_change_control"), heading)


def _add_commercial_proposal_section(
    doc: docx.Document, data: dict, heading: str = "12.0 Commercial Proposal & Total Cost of Ownership (TCO)",
) -> bool:
    try:
        if not any([
            data.get("pricing_approach"), data.get("pricing_table"),
            data.get("ongoing_costs_approach"), data.get("rate_card_and_change_control"),
        ]):
            return False

        doc.add_heading(heading, level=1)
        _add_pricing_approach_section(doc, data)
        _add_ongoing_costs_approach_section(doc, data)
        _add_rate_card_and_change_control_section(doc, data)
        return True
    except Exception:
        traceback.print_exc()
        return False


# ============================================================
# 13.0 OPEN QUESTIONS FOR CLARIFICATION
# ============================================================

def _add_open_questions_section(
    doc: docx.Document, questions, heading: str = "13.0 Open Questions for Clarification",
) -> bool:
    try:
        values = [str(v) for v in (questions or []) if isinstance(v, str) and v.strip()]
        if not values:
            return False
        doc.add_heading(heading, level=1)
        for v in values:
            doc.add_paragraph(v, style="List Bullet")
        return True
    except Exception:
        traceback.print_exc()
        return False


# ============================================================
# 14.0 UNADDRESSED ITEMS
# ============================================================

def _add_unaddressed_items_section(
    doc: docx.Document, items, heading: str = "14.0 Unaddressed Items",
) -> bool:
    """14.0 Unaddressed Items -- carried forward from the SOURCE
    requirements_summary.json's own inferred_requirements/coverage_gaps,
    explicitly flagged here rather than silently dropped."""
    try:
        values = [str(v) for v in (items or []) if isinstance(v, str) and v.strip()]
        if not values:
            return False
        doc.add_heading(heading, level=1)
        doc.add_paragraph(
            "The following items were flagged in the original requirements analysis as "
            "inferred or as coverage gaps, not as stated requirements -- they are noted here "
            "for completeness but are not addressed as formal compliance items above."
        )
        for v in values:
            doc.add_paragraph(v, style="List Bullet")
        return True
    except Exception:
        traceback.print_exc()
        return False
