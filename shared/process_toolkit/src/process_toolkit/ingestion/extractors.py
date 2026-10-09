"""Per-extension text extractors for load_directory_context.

Ported verbatim from the ADK sample's process_agents/common/utils.py
(_extract_*_file + _DIRECTORY_CONTEXT_EXTRACTORS) -- pure Python, nothing
framework-specific, so the logic carries over unchanged.
"""

import os


def _extract_txt_file(path: str) -> str:
    with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
        return f.read()


def _extract_docx_file(path: str) -> str:
    import docx
    doc = docx.Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pdf_file(path: str) -> str:
    # pypdf extracts embedded text only -- a scanned/image-only PDF with no
    # text layer yields empty strings per page (no OCR is attempted), which
    # load_directory_context surfaces as "no extractable text" rather than
    # silently pretending the file contributed content.
    from pypdf import PdfReader
    reader = PdfReader(path)
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(p for p in pages if p.strip())


def _extract_eml_file(path: str) -> str:
    import email
    from email import policy
    with open(path, "rb") as f:
        msg = email.message_from_binary_file(f, policy=policy.default)

    header_lines = [
        f"{header}: {msg.get(header)}" for header in ("Subject", "From", "To", "Date")
        if msg.get(header)
    ]

    body = ""
    if msg.is_multipart():
        # Prefer the first real text/plain body part; a part that also
        # carries a filename is an attachment, not the message body, even
        # if its content type happens to be text/plain (e.g. a .txt
        # attachment) -- skip those.
        for part in msg.walk():
            if part.get_content_type() == "text/plain" and not part.get_filename():
                body = part.get_content()
                break
    else:
        body = msg.get_content()

    return "\n".join(header_lines) + "\n\n" + (body or "")


def _extract_msg_file(path: str) -> str:
    import extract_msg
    msg = extract_msg.Message(path)
    try:
        header_lines = [
            f"{label}: {value}" for label, value in (
                ("Subject", msg.subject), ("From", msg.sender),
                ("To", msg.to), ("Date", msg.date),
            ) if value
        ]
        return "\n".join(header_lines) + "\n\n" + (msg.body or "")
    finally:
        msg.close()


_EXCEL_MAX_ROWS_PER_SHEET = 200


def _extract_excel_file(path: str) -> str:
    # .xls (legacy binary) needs xlrd, .xlsx (modern XML) needs openpyxl --
    # pandas picks the right engine automatically from the file extension,
    # as long as both packages are installed.
    import pandas as pd
    sheets = pd.read_excel(path, sheet_name=None)
    parts = []
    for sheet_name, df in sheets.items():
        truncated_df = df.head(_EXCEL_MAX_ROWS_PER_SHEET)
        note = ""
        if len(df) > _EXCEL_MAX_ROWS_PER_SHEET:
            note = f"\n... ({len(df) - _EXCEL_MAX_ROWS_PER_SHEET} more rows truncated)"
        parts.append(f"[Sheet: {sheet_name}]\n{truncated_df.to_string(index=False)}{note}")
    return "\n\n".join(parts)


# Legacy .doc (binary, pre-2007 Word) is deliberately absent -- python-docx
# only reads the modern .docx XML format. A .doc file is caught in
# load_directory_context and reported in files_skipped with a specific
# "re-save as .docx" reason, rather than silently falling through to
# "unsupported file type".
DIRECTORY_CONTEXT_EXTRACTORS = {
    ".txt": _extract_txt_file,
    ".md": _extract_txt_file,
    ".docx": _extract_docx_file,
    ".pdf": _extract_pdf_file,
    ".eml": _extract_eml_file,
    ".msg": _extract_msg_file,
    ".xls": _extract_excel_file,
    ".xlsx": _extract_excel_file,
}


def load_directory_context(directory: str, max_chars: int = 150000) -> dict:
    """
    Reads every supported file directly inside `directory` (not recursive --
    only that directory's own files, not subdirectories) and returns their
    extracted text, concatenated and labeled by filename, for use as source
    material in a requirements-extraction step. Turning the returned text
    into a structured requirements JSON is the calling agent's own job via
    its instructions.

    Returns {"status": "NOT_FOUND", "directory": <resolved path>} if the
    directory doesn't exist. Otherwise {"status": "OK", "directory": ...,
    "files_processed": [...], "files_skipped": [{"file":..., "reason":...}],
    "combined_text": "..."}. combined_text is capped at max_chars with a
    truncation note appended if exceeded, so a large directory can't blow
    out the model's context window.
    """
    resolved = os.path.abspath(directory)
    if not os.path.isdir(resolved):
        return {"status": "NOT_FOUND", "directory": resolved}

    try:
        entries = sorted(os.listdir(resolved))
    except Exception as e:
        return {"status": "ERROR", "directory": resolved, "error": str(e)}

    files_processed = []
    files_skipped = []
    sections = []

    for name in entries:
        full_path = os.path.join(resolved, name)
        if not os.path.isfile(full_path):
            continue

        ext = os.path.splitext(name)[1].lower()
        extractor = DIRECTORY_CONTEXT_EXTRACTORS.get(ext)
        if extractor is None:
            reason = (
                "legacy .doc (unsupported -- re-save as .docx)" if ext == ".doc"
                else f"unsupported file type '{ext or '(no extension)'}'"
            )
            files_skipped.append({"file": name, "reason": reason})
            continue

        try:
            text = (extractor(full_path) or "").strip()
        except Exception as e:
            files_skipped.append({"file": name, "reason": f"failed to read: {e}"})
            continue

        if not text:
            files_skipped.append({"file": name, "reason": "no extractable text"})
            continue

        files_processed.append(name)
        sections.append(f"=== {name} ===\n{text}")

    combined_text = "\n\n".join(sections)
    truncated = len(combined_text) > max_chars
    if truncated:
        combined_text = combined_text[:max_chars] + (
            f"\n\n[... truncated at {max_chars} characters ...]"
        )

    result = {
        "status": "OK",
        "directory": resolved,
        "files_processed": files_processed,
        "files_skipped": files_skipped,
        "combined_text": combined_text,
    }
    if truncated:
        result["truncated"] = True
    return result
