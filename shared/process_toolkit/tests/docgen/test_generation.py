"""End-to-end coverage for generation.py's own page-break bookkeeping --
NOT per-section content (see test_content.py/test_design_sections.py for
that), just "does the final document ever place two page breaks back to
back with nothing real between them," which Word renders as a fully
blank page.
"""

import json

import docx
from docx.oxml.ns import qn

from process_toolkit import paths
from process_toolkit.schema.process_json import PROCESS_JSON_FILENAME, detect_schema_type_from_disk
from process_toolkit.schema.requirements_json import REQUIREMENTS_SUMMARY_FILENAME
from process_toolkit.schema.proposal_response_json import PROPOSAL_RESPONSE_FILENAME
from process_toolkit.docgen.generation import create_standard_doc_from_file
from process_toolkit.docgen.structure import add_iso_page_break


def _has_page_break(p_elem) -> bool:
    return any(br.get(qn("w:type")) == "page" for br in p_elem.findall(".//" + qn("w:br")))


def test_add_iso_page_break_inserts_exactly_one_new_paragraph():
    # Regression test: an earlier version of this function ALSO inserted
    # a separate, empty "spacer" paragraph immediately before the real
    # page-break paragraph. When the PRECEDING section's own content
    # already filled the current page right to its bottom edge, that
    # spacer paragraph would itself overflow onto a fresh page (plain
    # text-wrap, no break needed for that alone), and the explicit break
    # right after it then forced ANOTHER page turn -- leaving one page
    # with nothing on it but that lone spacer. Confirmed directly
    # against a real generated design document built from a user's own
    # data: 3 of 19 forced breaks landed on exactly this edge case, each
    # rendering as a fully blank page. Asserting exactly one new
    # paragraph here is what rules out that double-page-turn.
    doc = docx.Document()
    doc.add_paragraph("Some content.")
    before = len(doc.paragraphs)

    add_iso_page_break(doc)

    after = doc.paragraphs
    assert len(after) == before + 1
    assert _has_page_break(after[-1]._p)


def _has_image(p_elem) -> bool:
    return len(p_elem.findall(".//" + qn("w:drawing"))) > 0


def _has_text(p_elem) -> bool:
    return any((t.text or "").strip() for t in p_elem.findall(".//" + qn("w:t")))


def _blank_page_gaps(docx_path: str) -> list:
    """
    Returns a list of (break_index_a, break_index_b) pairs for every
    place two page breaks appear with NO real content (text, image, or
    table) between them -- i.e. every FULLY BLANK page Word would
    actually render. Walks doc.element.body directly (not
    doc.paragraphs/doc.tables separately), since those two don't
    preserve true document order relative to each other -- a table
    sitting between two breaks would be invisible to a paragraphs-only
    check and wrongly reported as a blank page.
    """
    doc = docx.Document(docx_path)
    items = []
    for child in doc.element.body:
        if child.tag == qn("w:p"):
            items.append(("p", child))
        elif child.tag == qn("w:tbl"):
            items.append(("tbl", child))

    break_idxs = [i for i, (t, el) in enumerate(items) if t == "p" and _has_page_break(el)]
    gaps = []
    for bi in range(len(break_idxs) - 1):
        segment = items[break_idxs[bi] + 1:break_idxs[bi + 1]]
        has_real_content = any(
            kind == "tbl" or (kind == "p" and (_has_text(el) or _has_image(el)))
            for kind, el in segment
        )
        if not has_real_content:
            gaps.append((break_idxs[bi], break_idxs[bi + 1]))
    return gaps


# Deliberately has NO "appendix" key, and every other top-level key is one
# generation.py's own consumed_keys set already accounts for -- so BOTH
# _add_appendix_from_json (gated on "if appendix:") and
# _add_additional_data_section (nothing left over once every field is
# consumed) have nothing to render. This is the exact combination that
# triggered the bug: two unconditional add_iso_page_break(doc) calls
# fired back-to-back around these two no-op sections, rendering as a
# fully blank page right before "Appendix B" in a real generated
# document, confirmed directly.
SAMPLE_PROCESS_NO_APPENDIX_NO_EXTRA_DATA = {
    "process_name": "Vendor Onboarding",
    "version": "1.0",
    "industry_sector": "Procurement",
    "introduction": "Onboards new vendors with KYC checks.",
    "assumptions": ["Vendor has a valid tax ID."],
    "constraints": ["Must complete within 5 business days."],
    "stakeholders": [{"stakeholder_name": "Procurement Lead", "responsibilities": ["Approve vendor"]}],
    "process_steps": [
        {
            "step_name": "Collect Vendor Documents", "description": "Gather required docs.",
            "responsible_party": "Procurement Lead", "dependencies": [],
        },
    ],
    "governance_requirements": ["Quarterly audit"],
    "risks_and_controls": [{"risk": "Fraud", "control": "KYC check"}],
    "process_triggers": ["New vendor request submitted"],
    "process_end_conditions": ["Vendor approved and onboarded"],
    "change_management": [{"change_request_process": "Submit via ticket", "versioning_rules": "Semantic versioning"}],
    "continuous_improvement": [{"review_frequency": "Quarterly", "improvement_inputs": ["Audit findings"]}],
}


def test_no_blank_page_when_appendix_and_additional_data_are_both_absent():
    with open(paths.output_path(PROCESS_JSON_FILENAME), "w", encoding="utf-8") as f:
        json.dump(SAMPLE_PROCESS_NO_APPENDIX_NO_EXTRA_DATA, f)

    result = create_standard_doc_from_file("Vendor Onboarding", schema_type="process")
    assert result.startswith("SUCCESS:")

    out_path = result.split("SUCCESS: Professional document saved at ", 1)[1]
    gaps = _blank_page_gaps(out_path)
    assert gaps == [], f"found {len(gaps)} fully blank page(s) at break-index pairs {gaps}"


SAMPLE_REQUIREMENTS_SUMMARY = {
    "source_directory": "./vendor-docs",
    "source_files": ["rfp.pdf"],
    "summary": "Requirements for a vendor onboarding portal.",
    "industry_sector": "Retail",
    "stakeholders": ["Procurement Lead", "Legal"],
    "assumptions": ["Vendors already have a tax id."],
    "risks": ["Vendor data may be incomplete."],
    "open_questions": ["Which regions are in scope?"],
    "requirements_register": [
        {"id": "FR-001", "type": "functional", "description": "Allow vendor registration.",
         "source": "rfp.pdf", "acceptance_criteria": ["A vendor can submit a registration form."],
         "status": "proposed"},
        {"id": "NFR-001", "type": "non_functional", "description": "Respond within 2 seconds.",
         "source": "rfp.pdf", "acceptance_criteria": ["p95 latency < 2s"], "status": "proposed"},
    ],
    "requirement_priorities": [
        {"requirement_id": "FR-001", "requirement_text": "vendor registration", "priority": "high",
         "rationale": "core flow"},
    ],
    "conflicting_requirements": [],
    "requirements_needing_clarification": [],
    "inferred_requirements": [],
    "coverage_gaps": [{"category": "Data Privacy & Compliance", "description": "No retention policy stated."}],
}


class TestRequirementsDocumentGeneration:
    def test_generates_a_requirements_document(self):
        with open(paths.output_path(REQUIREMENTS_SUMMARY_FILENAME), "w", encoding="utf-8") as f:
            json.dump(SAMPLE_REQUIREMENTS_SUMMARY, f)

        result = create_standard_doc_from_file("Vendor Onboarding Requirements", schema_type="requirements")
        assert result.startswith("SUCCESS:")

        out_path = result.split("SUCCESS: Professional document saved at ", 1)[1]
        doc = docx.Document(out_path)
        headings = [p.text for p in doc.paragraphs if p.style and p.style.name == "Heading 1"]
        assert "1.0 Introduction and Purpose" in headings
        assert "2.0 Stakeholders" in headings
        assert "3.0 Requirements Register" in headings
        assert "4.0 Requirement Priorities" in headings
        # Conflicting/clarification/inferred are empty in this sample --
        # must not render a heading for a section with nothing to show.
        assert not any("Conflicting Requirements" in h for h in headings)
        assert not any("Needing Clarification" in h for h in headings)
        assert not any("Inferred Requirements" in h for h in headings)
        assert any("Coverage Gaps" in h for h in headings)
        # No diagram, no fabricated glossary.
        assert not any("Glossary" in h for h in headings)
        assert not any("Architecture Flow Diagram" in h for h in headings)

    def test_no_blank_pages_in_a_generated_requirements_document(self):
        with open(paths.output_path(REQUIREMENTS_SUMMARY_FILENAME), "w", encoding="utf-8") as f:
            json.dump(SAMPLE_REQUIREMENTS_SUMMARY, f)

        result = create_standard_doc_from_file("Vendor Onboarding Requirements", schema_type="requirements")
        out_path = result.split("SUCCESS: Professional document saved at ", 1)[1]
        gaps = _blank_page_gaps(out_path)
        assert gaps == [], f"found {len(gaps)} fully blank page(s) at break-index pairs {gaps}"

    def test_requirements_is_reachable_only_via_an_explicit_schema_type(self):
        # The one guarantee this whole feature depends on: a saved
        # requirements_summary.json must NEVER be silently auto-discovered
        # by detect_schema_type_from_disk()'s own fallback -- only an
        # explicitly-passed schema_type="requirements" reaches it.
        with open(paths.output_path(REQUIREMENTS_SUMMARY_FILENAME), "w", encoding="utf-8") as f:
            json.dump(SAMPLE_REQUIREMENTS_SUMMARY, f)

        assert detect_schema_type_from_disk() == "process"

        result = create_standard_doc_from_file("X", schema_type=None)
        assert result.startswith("ERROR:")
        assert PROCESS_JSON_FILENAME in result


SAMPLE_PROPOSAL_RESPONSE = {
    "summary": "Our response to the vendor onboarding portal RFP.",
    "our_understanding": "You need a self-service vendor registration portal with fast response times.",
    "proposed_solution": "A cloud-native portal built on our existing vendor management platform.",
    "architecture_components": [
        {"name": "Vendor Portal UI", "description": "Self-service registration frontend.",
         "connects_to": ["Vendor API"]},
        {"name": "Vendor API", "description": "REST API backing the portal.",
         "connects_to": ["ERP Connector"]},
        {"name": "ERP Connector", "description": "Syncs vendor records to the client's ERP.",
         "connects_to": []},
    ],
    "differentiators": ["10 years in vendor management SaaS.", "24/7 support included."],
    "compliance_matrix": [
        {"requirement_id": "FR-001", "requirement_text": "Allow vendor registration.",
         "requirement_type": "functional", "response_status": "Fully Compliant",
         "response": "Our platform already supports self-service vendor registration."},
        {"requirement_id": "NFR-001", "requirement_text": "Respond within 2 seconds.",
         "requirement_type": "non_functional", "response_status": "Partially Compliant",
         "response": "p95 latency is currently 2.3s; a caching layer upgrade will close this gap."},
    ],
    "assumptions": ["Client provides API access to their existing ERP."],
    "risks_and_mitigations": [{"risk": "Data migration delays.", "mitigation": "Phased cutover plan."}],
    "qualifications": "Our team has delivered 12 similar vendor portals in the retail sector.",
    "pricing_approach": "Fixed-price implementation plus a per-seat annual subscription.",
    "open_questions_for_client": ["Which regions are in scope for data residency?"],
    "unaddressed_items": ["Data retention policy was flagged as a gap in the original requirements analysis."],
}


class TestProposalResponseDocumentGeneration:
    def test_generates_a_proposal_response_document(self):
        with open(paths.output_path(PROPOSAL_RESPONSE_FILENAME), "w", encoding="utf-8") as f:
            json.dump(SAMPLE_PROPOSAL_RESPONSE, f)

        result = create_standard_doc_from_file("Vendor Portal Proposal Response", schema_type="proposal_response")
        assert result.startswith("SUCCESS:")

        out_path = result.split("SUCCESS: Professional document saved at ", 1)[1]
        doc = docx.Document(out_path)
        headings = [p.text for p in doc.paragraphs if p.style and p.style.name == "Heading 1"]
        assert "1.0 Executive Summary" in headings
        assert "2.0 Our Understanding of Your Requirements" in headings
        assert "3.0 Proposed Solution" in headings
        assert any("Solution Architecture" in h for h in headings)
        assert any("Differentiators" in h for h in headings)
        assert any("Compliance Matrix" in h for h in headings)
        assert any("Qualifications" in h for h in headings)
        assert any("Assumptions" in h for h in headings)
        assert any("Risks and Mitigations" in h for h in headings)
        assert any("Pricing Approach" in h for h in headings)
        assert any("Open Questions" in h for h in headings)
        assert any("Unaddressed Items" in h for h in headings)
        # The solution-architecture diagram embeds as a real inline image.
        assert len(doc.inline_shapes) >= 1
        # Still no fabricated glossary, and not the process/design-style
        # flow diagram heading (a different section entirely).
        assert not any("Glossary" in h for h in headings)
        assert not any("Architecture Flow Diagram" in h for h in headings)

    def test_no_solution_architecture_section_when_components_are_absent(self):
        # No-fabrication guarantee: a proposal response that never
        # populated architecture_components gets no diagram section at
        # all, not a placeholder/empty one.
        data = dict(SAMPLE_PROPOSAL_RESPONSE)
        data.pop("architecture_components")
        with open(paths.output_path(PROPOSAL_RESPONSE_FILENAME), "w", encoding="utf-8") as f:
            json.dump(data, f)

        result = create_standard_doc_from_file("Vendor Portal Proposal Response", schema_type="proposal_response")
        out_path = result.split("SUCCESS: Professional document saved at ", 1)[1]
        doc = docx.Document(out_path)
        headings = [p.text for p in doc.paragraphs if p.style and p.style.name == "Heading 1"]
        assert not any("Solution Architecture" in h for h in headings)
        assert len(doc.inline_shapes) == 0

    def test_no_blank_pages_in_a_generated_proposal_response_document(self):
        with open(paths.output_path(PROPOSAL_RESPONSE_FILENAME), "w", encoding="utf-8") as f:
            json.dump(SAMPLE_PROPOSAL_RESPONSE, f)

        result = create_standard_doc_from_file("Vendor Portal Proposal Response", schema_type="proposal_response")
        out_path = result.split("SUCCESS: Professional document saved at ", 1)[1]
        gaps = _blank_page_gaps(out_path)
        assert gaps == [], f"found {len(gaps)} fully blank page(s) at break-index pairs {gaps}"

    def test_proposal_response_is_reachable_only_via_an_explicit_schema_type(self):
        # Same guarantee as requirements: a saved proposal_response.json
        # must NEVER be silently auto-discovered by detect_schema_type_
        # from_disk()'s own fallback -- only an explicitly-passed
        # schema_type="proposal_response" reaches it.
        with open(paths.output_path(PROPOSAL_RESPONSE_FILENAME), "w", encoding="utf-8") as f:
            json.dump(SAMPLE_PROPOSAL_RESPONSE, f)

        assert detect_schema_type_from_disk() == "process"

        result = create_standard_doc_from_file("X", schema_type=None)
        assert result.startswith("ERROR:")
        assert PROCESS_JSON_FILENAME in result
