import docx

from process_toolkit.docgen import requirements_sections as rs


def test_overview_section_always_renders_even_with_no_data():
    doc = docx.Document()
    assert rs._add_requirements_overview_section(doc, {}) is True
    assert any(p.text == "1.0 Introduction and Purpose" for p in doc.paragraphs)


def test_overview_section_renders_assumptions_risks_and_open_questions_as_subsections():
    doc = docx.Document()
    data = {
        "summary": "A summary.",
        "assumptions": ["Vendors have a tax id."],
        "risks": ["Data may be incomplete."],
        "open_questions": ["Which regions are in scope?"],
        "industry_sector": "Retail",
        "source_directory": "./vendor-docs",
        "source_files": ["rfp.pdf"],
    }
    assert rs._add_requirements_overview_section(doc, data) is True
    headings = [p.text for p in doc.paragraphs if p.style and p.style.name == "Heading 2"]
    assert headings == ["1.1 Assumptions", "1.2 Risks", "1.3 Open Questions"]


def test_requirements_register_section_is_a_noop_with_nothing_present():
    doc = docx.Document()
    assert rs._add_requirements_register_section(doc, {}) is False


def test_requirements_register_section_splits_by_type():
    doc = docx.Document()
    data = {
        "requirements_register": [
            {"id": "FR-001", "type": "functional", "description": "Allow login.", "source": "a.txt",
             "acceptance_criteria": ["x"], "status": "proposed"},
            {"id": "NFR-001", "type": "non_functional", "description": "Fast.", "source": "a.txt",
             "acceptance_criteria": [], "status": "proposed"},
            {"id": "GOAL-001", "type": "goal", "description": "Grow revenue.", "source": "a.txt",
             "acceptance_criteria": [], "status": "proposed"},
            {"id": "CON-001", "type": "constraint", "description": "Must use existing ERP.",
             "source": "a.txt", "acceptance_criteria": [], "status": "proposed"},
        ],
    }
    assert rs._add_requirements_register_section(doc, data) is True
    headings = [p.text for p in doc.paragraphs if p.style and p.style.name == "Heading 2"]
    assert headings == [
        "3.1 Functional Requirements", "3.2 Non-Functional Requirements",
        "3.3 Goals and Objectives", "3.4 Constraints",
    ]
    assert len(doc.tables) == 4


def test_requirements_register_table_uses_na_for_a_missing_id():
    doc = docx.Document()
    rs._add_requirements_register_table(doc, [{"type": "functional", "description": "x"}])
    table = doc.tables[0]
    assert table.rows[1].cells[0].text == "N/A"


def test_old_format_summary_renders_plain_bullets_not_a_table():
    doc = docx.Document()
    data = {
        "functional_requirements": ["Allow login.", "Allow logout."],
        "non_functional_requirements": ["Must be fast."],
    }
    assert rs._add_requirements_register_section(doc, data) is True
    assert doc.tables == []
    bullets = [p.text for p in doc.paragraphs if p.style and p.style.name == "List Bullet"]
    assert bullets == ["Allow login.", "Allow logout.", "Must be fast."]


def test_requirement_priorities_section_sorts_by_priority_then_id():
    doc = docx.Document()
    priorities = [
        {"requirement_id": "FR-002", "requirement_text": "b", "priority": "low", "rationale": "r2"},
        {"requirement_id": "FR-001", "requirement_text": "a", "priority": "critical", "rationale": "r1"},
    ]
    assert rs._add_requirement_priorities_section(doc, priorities) is True
    table = doc.tables[0]
    assert table.rows[1].cells[0].text == "FR-001"
    assert table.rows[2].cells[0].text == "FR-002"


def test_requirement_priorities_section_is_a_noop_when_empty():
    doc = docx.Document()
    assert rs._add_requirement_priorities_section(doc, []) is False


def test_conflicting_requirements_section_joins_ids():
    doc = docx.Document()
    conflicts = [{"requirement_ids": ["FR-001", "FR-002"], "description": "d", "recommendation": "r"}]
    assert rs._add_conflicting_requirements_section(doc, conflicts) is True
    table = doc.tables[0]
    assert table.rows[1].cells[0].text == "FR-001, FR-002"


def test_requirements_needing_clarification_section_sorts_by_id():
    doc = docx.Document()
    items = [
        {"requirement_id": "FR-002", "requirement_text": "b", "question": "q2"},
        {"requirement_id": "FR-001", "requirement_text": "a", "question": "q1"},
    ]
    assert rs._add_requirements_needing_clarification_section(doc, items) is True
    table = doc.tables[0]
    assert table.rows[1].cells[0].text == "FR-001"


def test_inferred_requirements_section_includes_the_disclaimer():
    doc = docx.Document()
    items = [{"id": "INF-001", "requirement": "Audit logging.", "rationale": "compliance",
              "category": "non_functional", "derived_from": ["FR-001"], "acceptance_criteria": []}]
    assert rs._add_inferred_requirements_section(doc, items) is True
    assert any("NOT stated in the source material" in p.text for p in doc.paragraphs)
    table = doc.tables[0]
    assert table.rows[1].cells[4].text == "FR-001"


def test_coverage_gaps_section_sorts_alphabetically():
    doc = docx.Document()
    gaps = [{"category": "Security", "description": "d1"}, {"category": "Observability", "description": "d2"}]
    assert rs._add_coverage_gaps_section(doc, gaps) is True
    table = doc.tables[0]
    assert table.rows[1].cells[0].text == "Observability"
    assert table.rows[2].cells[0].text == "Security"


def test_all_section_renderers_are_noops_on_empty_input():
    doc = docx.Document()
    assert rs._add_requirement_priorities_section(doc, None) is False
    assert rs._add_conflicting_requirements_section(doc, None) is False
    assert rs._add_requirements_needing_clarification_section(doc, None) is False
    assert rs._add_inferred_requirements_section(doc, None) is False
    assert rs._add_coverage_gaps_section(doc, None) is False
    assert doc.tables == []
