import docx

from process_toolkit.docgen.content import _resolve_index_references, _add_process_steps_section


def _all_text(doc):
    return [p.text for p in doc.paragraphs]


class TestResolveIndexReferences:
    """_resolve_index_references -- a process step's own
    risks_and_controls/change_management/continuous_improvement
    sometimes comes back as bare integers referencing positions in the
    process-wide list of the same name, instead of the documented
    {"risk", "control"}-shaped (etc.) objects -- confirmed directly
    against a real generated document where every step rendered literal
    "0"/"2"/"4" bullets instead of any real risk/control text."""

    def test_resolves_valid_indices_to_real_objects(self):
        top_level = [{"risk": "A"}, {"risk": "B"}, {"risk": "C"}, {"risk": "D"}, {"risk": "E"}]
        assert _resolve_index_references([0, 2, 4], top_level) == [
            {"risk": "A"}, {"risk": "C"}, {"risk": "E"},
        ]

    def test_out_of_range_indices_are_dropped_not_rendered_as_content(self):
        top_level = [{"risk": "A"}, {"risk": "B"}]
        assert _resolve_index_references([0, 5, 1], top_level) == [{"risk": "A"}, {"risk": "B"}]

    def test_all_out_of_range_resolves_to_empty_not_the_bare_indices(self):
        top_level = [{"risk": "A"}]
        assert _resolve_index_references([5, 6], top_level) == []

    def test_no_top_level_list_resolves_to_empty_not_the_bare_indices(self):
        assert _resolve_index_references([0, 2], None) == []
        assert _resolve_index_references([0, 2], []) == []

    def test_well_formed_object_list_passes_through_unchanged(self):
        # The documented shape (list of {"risk", "control"} objects, not
        # bare integers) must never be altered by this guard.
        value = [{"risk": "A", "control": "X"}]
        assert _resolve_index_references(value, [{"risk": "unrelated"}]) is value

    def test_non_list_value_passes_through_unchanged(self):
        assert _resolve_index_references(None, [{"risk": "A"}]) is None
        assert _resolve_index_references("some string", [{"risk": "A"}]) == "some string"

    def test_booleans_are_not_treated_as_indices(self):
        # isinstance(True, int) is True in Python -- must not let a
        # stray boolean list slip through the "all bare ints" check.
        value = [True, False]
        assert _resolve_index_references(value, [{"risk": "A"}]) is value


class TestAddProcessStepsSectionResolvesIndexReferences:
    """End-to-end: a step whose risks_and_controls/change_management/
    continuous_improvement are index-reference lists (the real shape
    seen from a live generation run) must render the REAL referenced
    content, not bare numbers."""

    def test_index_reference_step_renders_real_risk_control_text(self):
        doc = docx.Document()
        steps = [{
            "step_name": "Screen Candidates",
            "description": "x" * 10,
            "step_risks_and_controls": [0, 2],
        }]
        top_level_risks_and_controls = [
            {"risk": "Legal liability from discrimination.", "control": "Merit-based criteria."},
            {"risk": "Unrelated risk.", "control": "Unrelated control."},
            {"risk": "Data breach.", "control": "Encryption at rest."},
        ]

        _add_process_steps_section(
            doc, steps, top_level_risks_and_controls=top_level_risks_and_controls,
        )

        text = _all_text(doc)
        assert any("Legal liability from discrimination." in t for t in text)
        assert any("Data breach." in t for t in text)
        # Index 1 wasn't referenced -- its content must not appear.
        assert not any("Unrelated risk." in t for t in text)
        # The bare index values themselves must never appear as bullets.
        assert not any(t.strip() in ("• 0", "• 2") for t in text)

    def test_step_with_only_out_of_range_indices_renders_no_section(self):
        doc = docx.Document()
        steps = [{
            "step_name": "Some Step",
            "description": "x" * 10,
            "change_management": [7],
        }]

        _add_process_steps_section(doc, steps, top_level_change_management=[{"change_request_process": "A"}])

        text = _all_text(doc)
        assert not any("Change Management" in t for t in text)

    def test_well_formed_step_level_content_still_renders_normally(self):
        # A step that already has REAL, well-formed risk/control objects
        # (not index references) must render exactly as before -- this
        # fix must not regress the already-working case.
        doc = docx.Document()
        steps = [{
            "step_name": "Some Step",
            "description": "x" * 10,
            "step_risks_and_controls": [{"risk": "Direct risk.", "control": "Direct control."}],
        }]

        _add_process_steps_section(doc, steps, top_level_risks_and_controls=[{"risk": "Unrelated."}])

        text = _all_text(doc)
        assert any("Direct risk." in t for t in text)
