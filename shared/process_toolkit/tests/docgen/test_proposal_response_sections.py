import docx

from process_toolkit.docgen import proposal_response_sections as prs


def test_executive_summary_always_renders_even_with_no_data():
    doc = docx.Document()
    assert prs._add_executive_summary_section(doc, {}) is True
    assert any(p.text == "1.0 Executive Summary" for p in doc.paragraphs)


def test_executive_summary_uses_the_provided_summary():
    doc = docx.Document()
    assert prs._add_executive_summary_section(doc, {"summary": "A tailored summary."}) is True
    assert any(p.text == "A tailored summary." for p in doc.paragraphs)


def test_executive_summary_renders_its_subsections_when_present():
    doc = docx.Document()
    data = {
        "our_understanding": "You need X.",
        "solution_summary": "A short solution summary.",
        "differentiators": ["Fast delivery."],
        "expected_outcomes": ["Reduced manual effort."],
    }
    assert prs._add_executive_summary_section(doc, data) is True
    headings = [p.text for p in doc.paragraphs if p.style and p.style.name == "Heading 2"]
    assert headings == [
        "1.1 Understanding of Client Objectives",
        "1.2 Summary of Proposed Solution",
        "1.3 Differentiators & Value Proposition",
        "1.4 Expected Outcomes & ROI",
    ]


def test_executive_summary_omits_subsection_headings_when_absent():
    doc = docx.Document()
    assert prs._add_executive_summary_section(doc, {}) is True
    assert not any(p.style and p.style.name == "Heading 2" for p in doc.paragraphs)


def test_understanding_section_is_a_noop_without_data():
    doc = docx.Document()
    assert prs._add_understanding_section(doc, {}) is False


def test_understanding_section_renders_when_present():
    doc = docx.Document()
    assert prs._add_understanding_section(doc, {"our_understanding": "You need X."}) is True
    assert any(p.text == "You need X." for p in doc.paragraphs)


def test_solution_summary_section_is_a_noop_without_data():
    doc = docx.Document()
    assert prs._add_solution_summary_section(doc, {}) is False


def test_expected_outcomes_section_renders_bullets():
    doc = docx.Document()
    assert prs._add_expected_outcomes_section(doc, ["Reduced manual effort."]) is True


def test_proposed_solution_section_is_a_noop_without_data():
    doc = docx.Document()
    assert prs._add_proposed_solution_section(doc, {}) is False


def test_solution_architecture_section_is_a_noop_when_empty():
    doc = docx.Document()
    assert prs._add_solution_architecture_section(doc, None, "Vendor Portal") is False
    assert prs._add_solution_architecture_section(doc, [{"description": "no name"}], "Vendor Portal") is False
    assert doc.tables == []


def test_solution_architecture_section_embeds_a_diagram_and_a_component_table():
    doc = docx.Document()
    components = [
        {"name": "Ingestion API", "description": "Receives documents.", "connects_to": ["OCR Engine"]},
        {"name": "OCR Engine", "description": "Extracts text.", "connects_to": []},
    ]
    assert prs._add_solution_architecture_section(doc, components, "Vendor Portal") is True
    assert any(p.text == "3.0 Solution Architecture" for p in doc.paragraphs)
    # A picture is embedded as an inline shape, not a paragraph/table.
    assert len(doc.inline_shapes) == 1
    table = doc.tables[0]
    assert table.rows[1].cells[0].text == "Ingestion API"
    assert table.rows[2].cells[0].text == "OCR Engine"


def test_scope_of_work_section_is_a_noop_when_every_subsection_is_empty():
    doc = docx.Document()
    assert prs._add_scope_of_work_section(doc, {}) is False
    assert not any(p.style and p.style.name == "Heading 1" for p in doc.paragraphs)


def test_scope_of_work_section_renders_present_subsections_only():
    doc = docx.Document()
    data = {"deliverables": ["A deployed pipeline.", "Admin dashboard."]}
    assert prs._add_scope_of_work_section(doc, data) is True
    headings = [p.text for p in doc.paragraphs if p.style and p.style.name in ("Heading 1", "Heading 2")]
    assert headings == ["4.0 Scope of Work & Boundaries", "4.4 Key Deliverables & Artifacts"]
    bullets = [p.text for p in doc.paragraphs if p.style and p.style.name == "List Bullet"]
    assert bullets == ["A deployed pipeline.", "Admin dashboard."]


def test_functional_coverage_section_computes_a_rollup_from_the_compliance_matrix():
    doc = docx.Document()
    matrix = [
        {"requirement_id": "FR-001", "response_status": "Fully Compliant"},
        {"requirement_id": "FR-002", "response_status": "Fully Compliant"},
        {"requirement_id": "FR-003", "response_status": "Partially Compliant"},
    ]
    assert prs._add_functional_coverage_section(doc, matrix) is True
    assert any("3 requirements" in p.text and "2 Fully Compliant" in p.text for p in doc.paragraphs)


def test_functional_coverage_section_is_a_noop_when_empty():
    doc = docx.Document()
    assert prs._add_functional_coverage_section(doc, None) is False


def test_differentiators_section_renders_bullets():
    doc = docx.Document()
    assert prs._add_differentiators_section(doc, ["Fast delivery.", "24/7 support."]) is True
    bullets = [p.text for p in doc.paragraphs if p.style and p.style.name == "List Bullet"]
    assert bullets == ["Fast delivery.", "24/7 support."]


def test_differentiators_section_is_a_noop_when_empty():
    doc = docx.Document()
    assert prs._add_differentiators_section(doc, None) is False


def test_compliance_matrix_section_sorts_by_requirement_id():
    doc = docx.Document()
    matrix = [
        {"requirement_id": "FR-002", "requirement_text": "b", "response_status": "Fully Compliant", "response": "r2"},
        {"requirement_id": "FR-001", "requirement_text": "a", "response_status": "Partially Compliant", "response": "r1"},
    ]
    assert prs._add_compliance_matrix_section(doc, matrix) is True
    table = doc.tables[0]
    assert table.rows[1].cells[0].text == "FR-001"
    assert table.rows[2].cells[0].text == "FR-002"


def test_compliance_matrix_section_renders_the_evidence_ref_column():
    doc = docx.Document()
    matrix = [
        {
            "requirement_id": "FR-001", "requirement_text": "a", "response_status": "Fully Compliant",
            "response": "r1", "evidence_ref": "See Proposed Solution",
        },
    ]
    assert prs._add_compliance_matrix_section(doc, matrix) is True
    table = doc.tables[0]
    assert table.rows[0].cells[4].text == "Evidence / Location"
    assert table.rows[1].cells[4].text == "• See Proposed Solution"


def test_compliance_matrix_section_is_a_noop_when_empty():
    doc = docx.Document()
    assert prs._add_compliance_matrix_section(doc, None) is False
    assert doc.tables == []


def test_delivery_approach_section_is_a_noop_when_every_subsection_is_empty():
    doc = docx.Document()
    assert prs._add_delivery_approach_section(doc, {}) is False


def test_delivery_approach_section_renders_present_subsections_only():
    doc = docx.Document()
    data = {"automation_approach": "We use an LLM-based test harness integrated with CI/CD."}
    assert prs._add_delivery_approach_section(doc, data) is True
    headings = [p.text for p in doc.paragraphs if p.style and p.style.name in ("Heading 1", "Heading 2")]
    assert headings == ["6.0 Delivery & Engagement Approach", "6.6 Automation & Tooling Approach"]


def test_slas_and_kpis_section_renders_a_table():
    doc = docx.Document()
    items = [{"metric": "Extraction accuracy", "target": ">= 92%"}]
    assert prs._add_slas_and_kpis_section(doc, items) is True
    table = doc.tables[0]
    assert table.rows[0].cells[0].text == "Metric"
    assert table.rows[1].cells[0].text == "Extraction accuracy"


def test_implementation_plan_section_is_a_noop_when_empty():
    doc = docx.Document()
    assert prs._add_implementation_plan_section(doc, None) is False
    assert doc.tables == []


def test_implementation_plan_section_renders_a_table_in_given_order():
    doc = docx.Document()
    plan = [
        {"phase": "Phase 2", "timing": "Weeks 5-8", "deliverable": "Integration.", "acceptance_criteria": "UAT passed."},
        {"phase": "Phase 1", "timing": "Weeks 1-4", "deliverable": "Design doc.", "acceptance_criteria": "Sign-off."},
    ]
    assert prs._add_implementation_plan_section(doc, plan) is True
    table = doc.tables[0]
    # Not sorted -- preserves the agent's own given (sequential) order.
    assert table.rows[1].cells[0].text == "Phase 2"
    assert table.rows[2].cells[0].text == "Phase 1"


def test_governance_section_is_a_noop_when_every_subsection_is_empty():
    doc = docx.Document()
    assert prs._add_governance_section(doc, {}) is False


def test_raci_matrix_section_renders_a_table():
    doc = docx.Document()
    items = [{"activity": "Sign-off", "responsible": "Delivery Lead", "accountable": "Client Sponsor",
              "consulted": "QA Lead", "informed": "Steering Committee"}]
    assert prs._add_raci_matrix_section(doc, items) is True
    table = doc.tables[0]
    assert table.rows[0].cells[0].text == "Activity"
    assert table.rows[1].cells[0].text == "• Sign-off"
    assert table.rows[1].cells[1].text == "Delivery Lead"


def test_governance_section_renders_present_subsections_only():
    doc = docx.Document()
    data = {"governance_framework": "A steering committee meets monthly."}
    assert prs._add_governance_section(doc, data) is True
    headings = [p.text for p in doc.paragraphs if p.style and p.style.name in ("Heading 1", "Heading 2")]
    assert headings == ["7.0 Governance Model, Roles & Responsibilities", "7.1 Governance Framework"]


def test_team_profiles_section_is_a_noop_when_every_subsection_is_empty():
    doc = docx.Document()
    assert prs._add_team_profiles_section(doc, {}) is False


def test_team_profiles_section_renders_present_subsections_only():
    doc = docx.Document()
    data = {"key_personnel": [{"role": "Delivery Lead", "responsibilities": "Overall delivery accountability."}]}
    assert prs._add_team_profiles_section(doc, data) is True
    headings = [p.text for p in doc.paragraphs if p.style and p.style.name in ("Heading 1", "Heading 2")]
    assert headings == ["8.0 Indicative Profiles & Team Structure", "8.2 Key Personnel Profiles"]
    table = doc.tables[0]
    assert table.rows[1].cells[0].text == "Delivery Lead"


def test_success_stories_section_is_a_noop_when_every_subsection_is_empty():
    doc = docx.Document()
    assert prs._add_success_stories_section(doc, {}) is False


def test_case_studies_section_renders_each_case_study_as_a_bold_title_plus_fields():
    doc = docx.Document()
    items = [{
        "title": "Enterprise Document Extraction", "context": "A regulated client.",
        "challenge": "Manual contract review.", "work_completed": "Deployed an AI pipeline.",
        "outcome": "50% faster processing.",
    }]
    assert prs._add_case_studies_section(doc, items) is True
    bold_runs = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "Enterprise Document Extraction" in bold_runs
    assert any("50% faster processing." in p.text for p in doc.paragraphs)


def test_case_studies_section_is_a_noop_without_a_title():
    doc = docx.Document()
    assert prs._add_case_studies_section(doc, [{"context": "no title"}]) is False


def test_qualifications_section_is_a_noop_without_data():
    doc = docx.Document()
    assert prs._add_qualifications_section(doc, {}) is False


def test_assumptions_dependencies_section_is_a_noop_when_every_subsection_is_empty():
    doc = docx.Document()
    assert prs._add_assumptions_dependencies_section(doc, {}) is False


def test_assumptions_section_renders_bullets():
    doc = docx.Document()
    assert prs._add_assumptions_section(doc, ["Client provides API access."]) is True


def test_client_dependencies_section_renders_bullets():
    doc = docx.Document()
    assert prs._add_client_dependencies_section(doc, ["Client provides SAP sandbox access."]) is True


def test_risks_and_mitigations_section_renders_a_table():
    doc = docx.Document()
    risks = [{"risk": "Scope creep.", "mitigation": "Change control process."}]
    assert prs._add_risks_and_mitigations_section(doc, risks) is True
    table = doc.tables[0]
    assert table.rows[1].cells[0].text == "• Scope creep."
    assert table.rows[1].cells[1].text == "• Change control process."


def test_assumptions_dependencies_section_renders_present_subsections_only():
    doc = docx.Document()
    data = {"assumptions": ["Client provides API access."]}
    assert prs._add_assumptions_dependencies_section(doc, data) is True
    headings = [p.text for p in doc.paragraphs if p.style and p.style.name in ("Heading 1", "Heading 2")]
    assert headings == ["11.0 Key Assumptions & Dependencies", "11.1 Project Assumptions"]


def test_pricing_approach_section_is_a_noop_without_data():
    doc = docx.Document()
    assert prs._add_pricing_approach_section(doc, {}) is False


def test_pricing_approach_section_renders_prose_only():
    doc = docx.Document()
    assert prs._add_pricing_approach_section(doc, {"pricing_approach": "Fixed price."}) is True
    assert any(p.text == "Fixed price." for p in doc.paragraphs)
    assert doc.tables == []


def test_pricing_approach_section_renders_the_pricing_table_without_prose():
    doc = docx.Document()
    data = {"pricing_table": [{"item": "Implementation", "basis": "One-time", "price": "TBD"}]}
    assert prs._add_pricing_approach_section(doc, data) is True
    table = doc.tables[0]
    assert table.rows[1].cells[0].text == "Implementation"
    assert table.rows[1].cells[2].text == "TBD"


def test_pricing_approach_section_defaults_a_missing_price_to_tbd():
    doc = docx.Document()
    data = {"pricing_table": [{"item": "Support", "basis": "Recurring"}]}
    assert prs._add_pricing_approach_section(doc, data) is True
    assert doc.tables[0].rows[1].cells[2].text == "TBD"


def test_commercial_proposal_section_is_a_noop_when_every_subsection_is_empty():
    doc = docx.Document()
    assert prs._add_commercial_proposal_section(doc, {}) is False


def test_commercial_proposal_section_renders_present_subsections_only():
    doc = docx.Document()
    data = {"rate_card_and_change_control": "Standard change-control process applies."}
    assert prs._add_commercial_proposal_section(doc, data) is True
    headings = [p.text for p in doc.paragraphs if p.style and p.style.name in ("Heading 1", "Heading 2")]
    assert headings == [
        "12.0 Commercial Proposal & Total Cost of Ownership (TCO)",
        "12.3 Rate Cards & Change Control Mechanics",
    ]


def test_open_questions_section_renders_bullets():
    doc = docx.Document()
    assert prs._add_open_questions_section(doc, ["Which region hosts the data?"]) is True


def test_unaddressed_items_section_includes_the_disclaimer():
    doc = docx.Document()
    assert prs._add_unaddressed_items_section(doc, ["Disaster recovery was flagged as a gap."]) is True
    assert any("inferred or as coverage gaps" in p.text for p in doc.paragraphs)


def test_all_section_renderers_are_noops_on_empty_input():
    doc = docx.Document()
    assert prs._add_understanding_section(doc, {}) is False
    assert prs._add_solution_summary_section(doc, {}) is False
    assert prs._add_differentiators_section(doc, None) is False
    assert prs._add_expected_outcomes_section(doc, None) is False
    assert prs._add_proposed_solution_section(doc, {}) is False
    assert prs._add_solution_architecture_section(doc, None, "Vendor Portal") is False
    assert prs._add_scope_of_work_section(doc, {}) is False
    assert prs._add_compliance_matrix_section(doc, None) is False
    assert prs._add_delivery_approach_section(doc, {}) is False
    assert prs._add_implementation_plan_section(doc, None) is False
    assert prs._add_governance_section(doc, {}) is False
    assert prs._add_team_profiles_section(doc, {}) is False
    assert prs._add_success_stories_section(doc, {}) is False
    assert prs._add_qualifications_section(doc, {}) is False
    assert prs._add_assumptions_dependencies_section(doc, {}) is False
    assert prs._add_assumptions_section(doc, None) is False
    assert prs._add_risks_and_mitigations_section(doc, None) is False
    assert prs._add_commercial_proposal_section(doc, {}) is False
    assert prs._add_pricing_approach_section(doc, {}) is False
    assert prs._add_open_questions_section(doc, None) is False
    assert prs._add_unaddressed_items_section(doc, None) is False
    assert doc.tables == []
