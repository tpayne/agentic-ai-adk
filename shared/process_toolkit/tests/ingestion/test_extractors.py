import os
import tempfile
from email.message import EmailMessage
from pathlib import Path

from process_toolkit.ingestion.extractors import (
    load_directory_context, DIRECTORY_CONTEXT_EXTRACTORS, _extract_msg_file,
    stage_attachments_to_directory,
)


def test_missing_directory_returns_not_found():
    with tempfile.TemporaryDirectory() as d:
        missing = str(Path(d) / "does-not-exist")
        result = load_directory_context(missing)
    assert result["status"] == "NOT_FOUND"


def test_resolving_to_agent_toolkit_root_returns_error_not_its_own_files():
    # AGENT_TOOL_PATH is always <project_root>/coded_tools (see cli.py and
    # every registry's own entry point) -- a directory argument that
    # resolves back to <project_root> itself (e.g. "." after the real path
    # was invalid) must be refused, not silently read as if it were genuine
    # user-supplied source material.
    with tempfile.TemporaryDirectory() as project_root:
        (Path(project_root) / "README.md").write_text("# Not real requirements source material")
        coded_tools_dir = Path(project_root) / "coded_tools"
        coded_tools_dir.mkdir()

        old_value = os.environ.get("AGENT_TOOL_PATH")
        os.environ["AGENT_TOOL_PATH"] = str(coded_tools_dir)
        try:
            result = load_directory_context(project_root)
        finally:
            if old_value is None:
                os.environ.pop("AGENT_TOOL_PATH", None)
            else:
                os.environ["AGENT_TOOL_PATH"] = old_value

    assert result["status"] == "ERROR"
    assert "files_processed" not in result
    assert "README.md" not in result.get("error", "")


def test_agent_tool_path_unset_does_not_affect_normal_directories():
    old_value = os.environ.pop("AGENT_TOOL_PATH", None)
    try:
        with tempfile.TemporaryDirectory() as d:
            (Path(d) / "notes.txt").write_text("Vendor onboarding requires KYC checks.")
            result = load_directory_context(d)
    finally:
        if old_value is not None:
            os.environ["AGENT_TOOL_PATH"] = old_value

    assert result["status"] == "OK"


def test_extracts_txt_and_docx_and_skips_legacy_doc():
    import docx

    with tempfile.TemporaryDirectory() as d:
        (Path(d) / "notes.txt").write_text(
            "Vendor onboarding requires KYC checks and a 3-day SLA."
        )

        doc = docx.Document()
        doc.add_paragraph("Stakeholders: Procurement Lead, Compliance Officer.")
        doc.save(str(Path(d) / "stakeholders.docx"))

        (Path(d) / "legacy.doc").write_text("ignored -- old binary format")

        result = load_directory_context(d)

    assert result["status"] == "OK"
    assert sorted(result["files_processed"]) == ["notes.txt", "stakeholders.docx"]
    assert "KYC checks" in result["combined_text"]
    assert "Procurement Lead" in result["combined_text"]
    assert "=== notes.txt ===" in result["combined_text"]

    skipped = {entry["file"]: entry["reason"] for entry in result["files_skipped"]}
    assert "legacy.doc" in skipped
    assert "re-save as .docx" in skipped["legacy.doc"]


def test_docx_table_text_is_extracted_not_silently_dropped():
    import docx

    with tempfile.TemporaryDirectory() as d:
        doc = docx.Document()
        doc.add_paragraph("Requirements register:")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "REQ-1"
        table.rows[0].cells[1].text = "Support SSO login"
        table.rows[1].cells[0].text = "REQ-2"
        table.rows[1].cells[1].text = "Encrypt data at rest"
        doc.save(str(Path(d) / "requirements.docx"))

        result = load_directory_context(d)

    assert result["status"] == "OK"
    assert "Support SSO login" in result["combined_text"]
    assert "Encrypt data at rest" in result["combined_text"]


def test_xml_file_is_extracted_as_raw_markup():
    with tempfile.TemporaryDirectory() as d:
        (Path(d) / "process.xml").write_text(
            "<process><step>Receive PO</step><step>Validate budget</step></process>"
        )

        result = load_directory_context(d)

    assert result["status"] == "OK"
    assert result["files_processed"] == ["process.xml"]
    assert "<step>Receive PO</step>" in result["combined_text"]


def test_unsupported_extension_is_skipped_not_fatal():
    with tempfile.TemporaryDirectory() as d:
        (Path(d) / "readme.txt").write_text("real content")
        (Path(d) / "image.jpg").write_bytes(b"\xff\xd8\xff\xe0not-a-real-jpeg")

        result = load_directory_context(d)

    assert result["status"] == "OK"
    assert result["files_processed"] == ["readme.txt"]
    skipped_files = [entry["file"] for entry in result["files_skipped"]]
    assert "image.jpg" in skipped_files


def test_excel_xlsx_renders_sheet_as_text():
    import pandas as pd

    with tempfile.TemporaryDirectory() as d:
        df = pd.DataFrame({"Step": ["Receive PO", "Validate budget"], "Owner": ["Ops", "Finance"]})
        xlsx_path = Path(d) / "steps.xlsx"
        df.to_excel(xlsx_path, index=False, sheet_name="Process")

        result = load_directory_context(d)

    assert result["status"] == "OK"
    assert result["files_processed"] == ["steps.xlsx"]
    assert "Receive PO" in result["combined_text"]
    assert "Validate budget" in result["combined_text"]
    assert "[Sheet: Process]" in result["combined_text"]


def test_eml_extraction_returns_subject_and_body():
    with tempfile.TemporaryDirectory() as d:
        msg = EmailMessage()
        msg["Subject"] = "Q3 onboarding requirements"
        msg["From"] = "ops@example.com"
        msg["To"] = "architecture@example.com"
        msg.set_content("All new vendors must complete KYC within 5 business days.")

        eml_path = Path(d) / "requirements.eml"
        eml_path.write_bytes(bytes(msg))

        result = load_directory_context(d)

    assert result["status"] == "OK"
    assert result["files_processed"] == ["requirements.eml"]
    assert "Q3 onboarding requirements" in result["combined_text"]
    assert "KYC within 5 business days" in result["combined_text"]


def test_blank_pdf_has_no_extractable_text_and_is_skipped():
    from pypdf import PdfWriter

    with tempfile.TemporaryDirectory() as d:
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        pdf_path = Path(d) / "blank.pdf"
        with open(pdf_path, "wb") as f:
            writer.write(f)

        result = load_directory_context(d)

    assert result["status"] == "OK"
    assert result["files_processed"] == []
    skipped = {entry["file"]: entry["reason"] for entry in result["files_skipped"]}
    assert skipped.get("blank.pdf") == "no extractable text"


def test_msg_extension_is_wired_to_its_extractor():
    # A full positive-content test would require fabricating a valid
    # OLE-format .msg file, which extract_msg (read-only) provides no
    # authoring API for. Covering the wiring and the graceful-failure
    # path (a malformed .msg must be skipped, not crash the directory
    # read) is the practical level of coverage here.
    assert DIRECTORY_CONTEXT_EXTRACTORS[".msg"] is _extract_msg_file

    with tempfile.TemporaryDirectory() as d:
        (Path(d) / "broken.msg").write_bytes(b"not a real OLE msg file")
        result = load_directory_context(d)

    assert result["status"] == "OK"
    assert result["files_processed"] == []
    skipped = {entry["file"]: entry["reason"] for entry in result["files_skipped"]}
    assert "broken.msg" in skipped
    assert "failed to read" in skipped["broken.msg"]


def test_combined_text_is_truncated_past_the_configured_limit():
    with tempfile.TemporaryDirectory() as d:
        (Path(d) / "big.txt").write_text("x" * 500)

        result = load_directory_context(d, max_chars=100)

    assert result["status"] == "OK"
    assert result.get("truncated") is True
    assert "truncated at 100 characters" in result["combined_text"]


class TestStageAttachmentsToDirectory:
    """stage_attachments_to_directory -- writes a chat turn's attached
    files to a fresh on-disk directory so they can be read through the
    exact same load_directory_context path a directory-context request
    already uses, instead of a second, parallel extraction mechanism."""

    def test_no_files_returns_none(self):
        assert stage_attachments_to_directory([]) == {"status": "NONE"}
        assert stage_attachments_to_directory(None) == {"status": "NONE"}

    def test_writes_files_into_a_fresh_directory_under_base_dir(self):
        with tempfile.TemporaryDirectory() as base_dir:
            result = stage_attachments_to_directory(
                [{"name": "notes.txt", "data": b"Vendor onboarding requires a 3-day SLA."}],
                base_dir=base_dir,
            )

        assert result["status"] == "OK"
        assert result["files_written"] == ["notes.txt"]
        # The staging directory itself is created fresh under base_dir --
        # not base_dir reused directly -- so concurrent turns never collide.
        assert Path(result["directory"]).parent == Path(base_dir)

    def test_staged_file_round_trips_through_load_directory_context(self):
        # The whole point of staging to disk is that the SAME
        # load_directory_context tool an explicit-directory request uses
        # can read it back -- exercise that full round trip, not just the
        # staging step in isolation.
        with tempfile.TemporaryDirectory() as base_dir:
            stage_result = stage_attachments_to_directory(
                [{"name": "notes.txt", "data": b"Vendor onboarding requires KYC within 5 days."}],
                base_dir=base_dir,
            )
            read_result = load_directory_context(stage_result["directory"])

        assert read_result["status"] == "OK"
        assert read_result["files_processed"] == ["notes.txt"]
        assert "KYC within 5 days" in read_result["combined_text"]

    def test_staged_docx_and_xml_attachments_round_trip_through_load_directory_context(self):
        # Regression test: an uploaded .docx or .xml file must be
        # extractable the same way a .txt attachment already is -- .docx
        # extraction already existed server-side (only the web client's
        # own <input accept> list blocked the upload before it ever
        # reached here); .xml extraction is new.
        import docx

        with tempfile.TemporaryDirectory() as source_dir:
            doc = docx.Document()
            doc.add_paragraph("Stakeholders: Procurement Lead, Compliance Officer.")
            docx_path = Path(source_dir) / "stakeholders.docx"
            doc.save(str(docx_path))
            docx_bytes = docx_path.read_bytes()

        with tempfile.TemporaryDirectory() as base_dir:
            stage_result = stage_attachments_to_directory(
                [
                    {"name": "stakeholders.docx", "data": docx_bytes},
                    {"name": "process.xml", "data": b"<process><step>Receive PO</step></process>"},
                ],
                base_dir=base_dir,
            )
            read_result = load_directory_context(stage_result["directory"])

        assert read_result["status"] == "OK"
        assert sorted(read_result["files_processed"]) == ["process.xml", "stakeholders.docx"]
        assert "Procurement Lead" in read_result["combined_text"]
        assert "<step>Receive PO</step>" in read_result["combined_text"]

    def test_multiple_files_are_all_staged(self):
        with tempfile.TemporaryDirectory() as base_dir:
            result = stage_attachments_to_directory(
                [
                    {"name": "a.txt", "data": b"first"},
                    {"name": "b.txt", "data": b"second"},
                ],
                base_dir=base_dir,
            )

        assert result["status"] == "OK"
        assert sorted(result["files_written"]) == ["a.txt", "b.txt"]

    def test_path_traversal_filename_is_reduced_to_its_safe_basename(self):
        # os.path.basename strips the directory components entirely --
        # "../../etc/cron.d/evil" becomes just "evil", written safely
        # inside the staging directory, never escaping it.
        with tempfile.TemporaryDirectory() as base_dir:
            result = stage_attachments_to_directory(
                [{"name": "../../etc/cron.d/evil", "data": b"whatever"}],
                base_dir=base_dir,
            )

            assert result["status"] == "OK"
            assert result["files_written"] == ["evil"]
            assert (Path(result["directory"]) / "evil").read_bytes() == b"whatever"
        assert not (Path(base_dir).parent / "etc").exists()

    def test_bare_dotdot_filename_is_rejected(self):
        with tempfile.TemporaryDirectory() as base_dir:
            result = stage_attachments_to_directory(
                [{"name": "..", "data": b"whatever"}],
                base_dir=base_dir,
            )
        assert result["status"] == "ERROR"

    def test_missing_base_dir_is_an_error_not_an_exception(self):
        with tempfile.TemporaryDirectory() as base_dir:
            missing = str(Path(base_dir) / "does-not-exist")
            result = stage_attachments_to_directory(
                [{"name": "a.txt", "data": b"x"}], base_dir=missing,
            )
        assert result["status"] == "ERROR"
        assert "error" in result

    def test_explicit_directory_is_used_exactly_and_created_if_missing(self):
        with tempfile.TemporaryDirectory() as base_dir:
            target = str(Path(base_dir) / "fixed-session-dir")
            result = stage_attachments_to_directory(
                [{"name": "a.txt", "data": b"x"}], directory=target,
            )
        assert result["status"] == "OK"
        assert result["directory"] == target

    def test_explicit_directory_clears_stale_files_from_an_earlier_turn(self):
        # A deterministic, reused-across-turns directory must never let a
        # later turn's attachment set mix with an earlier turn's leftovers.
        with tempfile.TemporaryDirectory() as base_dir:
            target = str(Path(base_dir) / "fixed-session-dir")
            first = stage_attachments_to_directory(
                [{"name": "first.txt", "data": b"first"}], directory=target,
            )
            second = stage_attachments_to_directory(
                [{"name": "second.txt", "data": b"second"}], directory=target,
            )

        assert first["status"] == "OK"
        assert second["status"] == "OK"
        assert second["files_written"] == ["second.txt"]
        assert not (Path(target) / "first.txt").exists()

    def test_explicit_directory_does_not_disturb_an_existing_subdirectory(self):
        # Clearing stale FILES must not reach into subdirectories --
        # confirms the clear step only removes files, not recursively.
        with tempfile.TemporaryDirectory() as base_dir:
            target = Path(base_dir) / "fixed-session-dir"
            target.mkdir()
            (target / "keep-me").mkdir()
            (target / "keep-me" / "nested.txt").write_text("should survive")

            result = stage_attachments_to_directory(
                [{"name": "a.txt", "data": b"x"}], directory=str(target),
            )

            assert result["status"] == "OK"
            assert (target / "keep-me" / "nested.txt").read_text() == "should survive"
